"""现场符合性巡查 v2：巡查任务、巡查记录、现场媒体、问题隐患与整改闭环。

任务驱动：平台派发任务并指定巡查员账号；巡查员在小程序登录若依后，
经 Java 代理调用本模块上传媒体、追加巡查记录、记录隐患；平台记录隐患、
下发整改要求并复核闭环。任务按指派账号做行级隔离。

认证：平台/Java 代理走 X-Service-Token（沿用 audit_api 既有约定）；
行级隔离所需的操作者身份由 Java 通过 X-Actor-* 头注入（Java 已完成
@PreAuthorize 角色校验，本模块据此过滤数据）。
"""

from __future__ import annotations

import hmac
import io
import json
import logging
import re
import sqlite3
import threading
import uuid
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import Annotated, Any, Iterator, Literal
from urllib.parse import quote

from fastapi import APIRouter, Body, Depends, File, Form, Header, HTTPException, UploadFile, status
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field

from .config import (
    PATROL_DB,
    PATROL_DEV_TOKEN,
    PATROL_UPLOAD_ROOT,
    PROJECT_ARCHIVE_DB,
    SERVICE_TOKEN,
)

LOGGER = logging.getLogger(__name__)

TASK_STATUSES = {"pending", "executing", "completed", "closed"}
_FIELD_CN = {"name": "任务名称", "lines": "线路", "user_ids": "指派巡查员", "location_desc": "位置", "requirement": "巡查内容"}
HAZARD_STATUSES = {"pending_confirm", "pending_rectify", "rectifying", "pending_review", "closed"}
RECORD_TYPES = {"patrol", "rectify"}
MEDIA_KINDS = {"photo", "video"}
DICT_TYPES = {"line", "construction_type", "hazard_type", "hazard_risk"}
PHOTO_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}
VIDEO_SUFFIXES = {".mp4", ".mov"}
MAX_PHOTOS_PER_RECORD = 9
MAX_VIDEOS_PER_RECORD = 2
MAX_PHOTO_BYTES = 20 * 1024 * 1024
MAX_VIDEO_BYTES = 50 * 1024 * 1024
# 审核意见 → 现场核查（Step①）
OPINION_STATUSES = {"pending_photo", "photo_taken", "done", "returned"}
MAX_OPINION_PHOTOS = 9
MAX_OPINION_PHOTO_BYTES = 20 * 1024 * 1024
# 监测方案文档
DOC_PDF_SUFFIXES = {".pdf"}
DOC_WORD_SUFFIXES = {".doc", ".docx"}
DOC_IMAGE_SUFFIXES = set(PHOTO_SUFFIXES)
DOC_SUFFIXES = DOC_PDF_SUFFIXES | DOC_WORD_SUFFIXES | DOC_IMAGE_SUFFIXES
MAX_DOC_BYTES = 50 * 1024 * 1024
MAX_DOCS_PER_TASK = 9


def _now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _today() -> str:
    return datetime.now().strftime("%Y%m%d")


def _id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex}"


def _text(value: Any, *, field: str, required: bool = False, maximum: int = 500) -> str:
    result = " ".join(str(value or "").split())
    if required and not result:
        raise ValueError(f"{field}不能为空。")
    if len(result) > maximum:
        raise ValueError(f"{field}不能超过{maximum}个字符。")
    return result


def _optional_float(value: Any, *, field: str, minimum: float, maximum: float) -> float | None:
    if value in (None, ""):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field}必须是数字。") from exc
    if not minimum <= result <= maximum:
        raise ValueError(f"{field}必须在{minimum}至{maximum}之间。")
    return result


def _optional_int(value: Any, *, field: str, minimum: int, maximum: int) -> int | None:
    if value is None:
        return None
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field}必须是整数。") from exc
    if not minimum <= result <= maximum:
        raise ValueError(f"{field}必须在{minimum}至{maximum}之间。")
    return result


def _limit_text(value: str, maximum: int) -> str:
    text = " ".join(str(value or "").split())
    return text if len(text) <= maximum else text[: maximum - 1].rstrip() + "…"


def _as_list(value: Any) -> list[str]:
    """把 数组 / 逗号顿号分隔串 / 单个值 归一成字符串列表。"""
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        raw = [str(v).strip() for v in value]
    else:
        raw = [p for p in re.split(r"[,，、;；\n\r]+", str(value)) if p.strip()]
    return [p for p in raw if p]


def _workflow_patrol_opinion(data: dict[str, Any]) -> str:
    """Turn the structured final audit result into readable field-inspection instructions."""
    raw = data.get("latest_result_json")
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except json.JSONDecodeError:
            raw = {}
    payload = raw if isinstance(raw, dict) else {}
    overall = payload.get("overall_opinion") or payload.get("overallOpinion") or {}
    if not isinstance(overall, dict):
        overall = {}
    items = payload.get("review_items") or payload.get("items") or []
    if not isinstance(items, list):
        items = []

    lines: list[str] = []
    conclusion = _limit_text(str(overall.get("conclusion") or data.get("latest_summary") or ""), 700)
    recommendation = _limit_text(str(overall.get("recommendation") or ""), 700)
    if conclusion:
        lines.append(f"终审综合结论：{conclusion}")
    if recommendation and recommendation != conclusion:
        lines.append(f"综合处置要求：{recommendation}")
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            continue
        title = _limit_text(str(item.get("title") or item.get("topic") or f"审核事项{index}"), 120)
        opinion = _limit_text(str(item.get("conclusion") or item.get("recommendation") or item.get("opinion") or ""), 600)
        if opinion:
            lines.append(f"{index}. {title}：{opinion}")
    final_approval = _limit_text(str(data.get("final_approval_opinion") or ""), 500)
    if final_approval and final_approval not in " ".join(lines):
        lines.append(f"终审人意见：{final_approval}")
    return _limit_text("；".join(lines) or "终审通过。请结合审核结论开展现场符合性巡查。", 1900)


def _workflow_overall_meta(data: dict[str, Any]) -> dict[str, str]:
    """从终审载荷提取“综合评价/终审人意见”文本（核查栏顶部综合评价卡用，不再拼进巡查内容）。"""
    raw = data.get("latest_result_json")
    if isinstance(raw, str) and raw.strip():
        try:
            raw = json.loads(raw)
        except json.JSONDecodeError:
            raw = {}
    payload = raw if isinstance(raw, dict) else {}
    overall = payload.get("overall_opinion") or payload.get("overallOpinion") or {}
    if not isinstance(overall, dict):
        overall = {}
    conclusion = str(overall.get("conclusion") or data.get("latest_summary") or "").strip()
    recommendation = str(overall.get("recommendation") or "").strip()
    overall_text = _limit_text(conclusion, 1600)
    if recommendation and recommendation not in (conclusion or ""):
        recommendation_block = f"综合处置要求：{recommendation}"
        overall_text = _limit_text(
            f"{conclusion}\n{recommendation_block}" if conclusion else recommendation_block,
            1900,
        )
    final_opinion = _limit_text(str(data.get("final_approval_opinion") or "").strip(), 800)
    return {"overall_conclusion": overall_text, "final_opinion": final_opinion}


def _patrol_requirement_text(opinion_count: int) -> str:
    """D120：巡查内容固定为“指引 + 意见条数”，不再拼接整段终审长文。"""
    return f"终审通过，共 {int(opinion_count or 0)} 条意见需现场核查，请逐条核实并拍照留证。"


# ---------------------------------------------------------------------------
# 审核意见 → 现场核查：从项目档案收集意见（解耦：只读项目档案 SQLite）
# ---------------------------------------------------------------------------

def _opinion_result_items(result_data: Any) -> list[tuple[str, str, str]]:
    """宽容提取审核结果里的逐条意见，返回 [(title, risk, content)]（去重、过滤已符合项）。"""
    payload = result_data if isinstance(result_data, dict) else {}
    items: list[Any] = []
    for key in ("review_items", "items", "opinions"):
        value = payload.get(key)
        if isinstance(value, list):
            items.extend(item for item in value if isinstance(item, dict))
    if not items:
        risk_report = ((payload.get("dynamic_regulation_audit") or {}).get("risk_report") or {})
        items = [item for item in (risk_report.get("findings") or []) if isinstance(item, dict)]
    if not items:
        details = payload.get("audit_details") or {}
        items = [item for item in (details.get("non_compliant_items") or []) if isinstance(item, dict)]

    seen: set[str] = set()
    out: list[tuple[str, str, str]] = []
    for item in items:
        judgement = str(item.get("judgement") or item.get("status") or "")
        if judgement == "compliant":
            continue
        title = str(item.get("title") or item.get("topic") or item.get("name") or "").strip()
        risk = str(item.get("risk_level") or item.get("severity") or "").strip()
        body = str(item.get("conclusion") or item.get("analysis") or item.get("result")
                   or item.get("recommendation") or item.get("opinion") or "").strip()
        recommendation = str(item.get("recommendation") or item.get("opinion") or "").strip()
        if recommendation and recommendation not in body:
            body = f"{body}；{recommendation}" if body else recommendation
        content = _limit_text(body or title, 1600)
        key = f"{title}\u0000{content}"[:400]
        if not content or key in seen:
            continue
        seen.add(key)
        out.append((_limit_text(title, 160), _limit_text(risk, 20), content))
    return out


def _resolve_archive_stage(project_id: str, stage_id: str, fallback_name: str) -> tuple[str, str]:
    """用项目档案里的“真实阶段”校正意见的 stage_id/stage_name。

    优先级：①已带 stage_id → 按 id 查档案现名（含改名）；②未带 id 且项目仅一个阶段 → 用该阶段；
    ③未带 id 且多个阶段 → 名称精确匹配；④都无法确认 → 保留流程快照值。
    """
    if not project_id:
        return stage_id or "", fallback_name
    db = Path(PROJECT_ARCHIVE_DB)
    if not db.exists():
        return stage_id or "", fallback_name
    try:
        connection = sqlite3.connect(db, timeout=10)
        try:
            connection.row_factory = sqlite3.Row
            rows = connection.execute(
                "SELECT stage_id, name FROM project_stages WHERE project_id = ? ORDER BY stage_order, stage_id",
                (project_id,),
            ).fetchall()
        finally:
            connection.close()
    except sqlite3.Error:
        LOGGER.exception("读取档案阶段失败：%s", project_id)
        return stage_id or "", fallback_name
    if not rows:
        return stage_id or "", fallback_name
    if stage_id:
        found = next((row for row in rows if row["stage_id"] == stage_id), None)
        if found is not None:
            return stage_id, found["name"] or fallback_name
        return stage_id, fallback_name
    if len(rows) == 1:
        return rows[0]["stage_id"], rows[0]["name"] or fallback_name
    exact_name = str(fallback_name or "").strip()
    if exact_name:
        matched = next((row for row in rows if row["name"] == exact_name), None)
        if matched is not None:
            return matched["stage_id"], matched["name"]
    return "", fallback_name


def _collect_project_opinions(project_id: str) -> list[dict[str, Any]]:
    """读取项目档案库中该项目已成功审核的 audit_records，提取逐条意见（带阶段上下文）。"""
    db = Path(PROJECT_ARCHIVE_DB)
    if not db.exists():
        return []
    rows: list[dict[str, Any]] = []
    try:
        connection = sqlite3.connect(db, timeout=10)
        try:
            connection.row_factory = sqlite3.Row
            records = connection.execute(
                """
                SELECT ar.stage_id, ar.result_json, ps.name AS stage_name
                FROM audit_records ar
                JOIN project_stages ps ON ps.stage_id = ar.stage_id
                WHERE ar.project_id = ? AND ar.status = 'success'
                ORDER BY ar.completed_at, ps.stage_order
                """,
                (project_id,),
            ).fetchall()
            for record in records:
                result_data: dict[str, Any] = {}
                raw = record["result_json"]
                if isinstance(raw, str) and raw.strip():
                    try:
                        result_data = json.loads(raw)
                    except json.JSONDecodeError:
                        result_data = {}
                for index, (title, risk, content) in enumerate(
                    _opinion_result_items(result_data), start=1
                ):
                    rows.append({
                        "stage_id": record["stage_id"] or "",
                        "stage_name": record["stage_name"] or "",
                        "opinion_no": index,
                        "title": title,
                        "risk_level": risk,
                        "content": content,
                    })
        finally:
            connection.close()
    except sqlite3.Error:
        LOGGER.exception("读取项目档案意见失败：%s", project_id)
    return rows


class PatrolRepository:
    """巡查任务、巡查记录、媒体、隐患与字典的持久化存储。"""

    def __init__(self, db_path: str | Path | None = None, upload_root: str | Path | None = None) -> None:
        self.db_path = Path(db_path or PATROL_DB).resolve()
        self.upload_root = Path(upload_root or PATROL_UPLOAD_ROOT).resolve()
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.upload_root.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self._initialize()

    @contextmanager
    def _connect(self) -> Iterator[sqlite3.Connection]:
        connection = sqlite3.connect(self.db_path, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA busy_timeout = 30000")
        try:
            with connection:
                yield connection
        finally:
            connection.close()

    def _initialize(self) -> None:
        with self._lock, self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS patrol_dict (
                    dict_id TEXT PRIMARY KEY,
                    dict_type TEXT NOT NULL,
                    label TEXT NOT NULL,
                    value TEXT NOT NULL DEFAULT '',
                    sort INTEGER NOT NULL DEFAULT 0,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS patrol_tasks (
                    task_id TEXT PRIMARY KEY,
                    task_no TEXT NOT NULL UNIQUE,
                    task_type TEXT NOT NULL DEFAULT 'construction'
                        CHECK(task_type IN ('pre_construction','construction')),
                    source_workflow_id TEXT NOT NULL DEFAULT '',
                    source_project_id TEXT NOT NULL DEFAULT '',
                    source_project_name TEXT NOT NULL DEFAULT '',
                    name TEXT NOT NULL DEFAULT '',
                    line TEXT NOT NULL DEFAULT '',
                    location_desc TEXT NOT NULL DEFAULT '',
                    requirement TEXT NOT NULL DEFAULT '',
                    assigned_user_id TEXT NOT NULL DEFAULT '',
                    assigned_user_name TEXT NOT NULL DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'pending'
                        CHECK(status IN ('pending','executing','completed','closed')),
                    dispatcher TEXT NOT NULL DEFAULT '',
                    dispatch_time TEXT NOT NULL DEFAULT '',
                    remark TEXT NOT NULL DEFAULT '',
                    monitor_frequency TEXT NOT NULL DEFAULT '',
                    monitor_points TEXT NOT NULL DEFAULT '',
                    warning_threshold TEXT NOT NULL DEFAULT '',
                    emergency_plan TEXT NOT NULL DEFAULT '',
                    report_requirement TEXT NOT NULL DEFAULT '',
                    review_opinion TEXT NOT NULL DEFAULT '',
                    legacy INTEGER NOT NULL DEFAULT 0,
                    deleted INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS patrol_records (
                    record_id TEXT PRIMARY KEY,
                    task_id TEXT NOT NULL,
                    type TEXT NOT NULL DEFAULT 'patrol'
                        CHECK(type IN ('patrol','rectify')),
                    hazard_id TEXT NOT NULL DEFAULT '',
                    title TEXT NOT NULL DEFAULT '',
                    longitude REAL,
                    latitude REAL,
                    accuracy REAL,
                    note TEXT NOT NULL DEFAULT '',
                    created_by TEXT NOT NULL DEFAULT '',
                    created_by_name TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(task_id) REFERENCES patrol_tasks(task_id) ON DELETE CASCADE
                );

                CREATE TABLE IF NOT EXISTS patrol_media (
                    media_id TEXT PRIMARY KEY,
                    record_id TEXT NOT NULL,
                    kind TEXT NOT NULL CHECK(kind IN ('photo','video')),
                    file_name TEXT NOT NULL DEFAULT '',
                    file_path TEXT NOT NULL DEFAULT '',
                    caption TEXT NOT NULL DEFAULT '',
                    taken_at TEXT NOT NULL DEFAULT '',
                    sort INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(record_id) REFERENCES patrol_records(record_id) ON DELETE CASCADE
                );

                CREATE TABLE IF NOT EXISTS patrol_hazards (
                    hazard_id TEXT PRIMARY KEY,
                    task_id TEXT NOT NULL,
                    record_id TEXT NOT NULL DEFAULT '',
                    media_id TEXT NOT NULL DEFAULT '',
                    video_time TEXT NOT NULL DEFAULT '',
                    description TEXT NOT NULL DEFAULT '',
                    hazard_type TEXT NOT NULL DEFAULT '',
                    risk_level TEXT NOT NULL DEFAULT '',
                    rectify_requirement TEXT NOT NULL DEFAULT '',
                    rectify_owner TEXT NOT NULL DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'pending_rectify'
                        CHECK(status IN ('pending_confirm','pending_rectify','rectifying','pending_review','closed')),
                    created_by TEXT NOT NULL DEFAULT '',
                    created_by_name TEXT NOT NULL DEFAULT '',
                    created_by_role TEXT NOT NULL DEFAULT 'platform',
                    confirmer TEXT NOT NULL DEFAULT '',
                    confirm_time TEXT NOT NULL DEFAULT '',
                    submitter TEXT NOT NULL DEFAULT '',
                    submit_time TEXT NOT NULL DEFAULT '',
                    reviewer TEXT NOT NULL DEFAULT '',
                    review_comment TEXT NOT NULL DEFAULT '',
                    review_time TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    FOREIGN KEY(task_id) REFERENCES patrol_tasks(task_id) ON DELETE CASCADE
                );

                CREATE TABLE IF NOT EXISTS patrol_hazard_shots (
                    shot_id TEXT PRIMARY KEY,
                    hazard_id TEXT NOT NULL,
                    file_name TEXT NOT NULL DEFAULT '',
                    file_path TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(hazard_id) REFERENCES patrol_hazards(hazard_id) ON DELETE CASCADE
                );

                CREATE TABLE IF NOT EXISTS patrol_task_docs (
                    doc_id TEXT PRIMARY KEY,
                    task_id TEXT NOT NULL,
                    file_name TEXT NOT NULL DEFAULT '',
                    file_path TEXT NOT NULL DEFAULT '',
                    kind TEXT NOT NULL DEFAULT 'pdf'
                        CHECK(kind IN ('pdf','word','image')),
                    size INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(task_id) REFERENCES patrol_tasks(task_id) ON DELETE CASCADE
                );

                CREATE TABLE IF NOT EXISTS patrol_task_lines (
                    task_id TEXT NOT NULL,
                    line TEXT NOT NULL,
                    sort INTEGER NOT NULL DEFAULT 0,
                    PRIMARY KEY(task_id, line),
                    FOREIGN KEY(task_id) REFERENCES patrol_tasks(task_id) ON DELETE CASCADE
                );

                CREATE TABLE IF NOT EXISTS patrol_task_users (
                    task_id TEXT NOT NULL,
                    user_id TEXT NOT NULL,
                    user_name TEXT NOT NULL DEFAULT '',
                    sort INTEGER NOT NULL DEFAULT 0,
                    PRIMARY KEY(task_id, user_id),
                    FOREIGN KEY(task_id) REFERENCES patrol_tasks(task_id) ON DELETE CASCADE
                );

                CREATE TABLE IF NOT EXISTS patrol_task_edit_logs (
                    log_id TEXT PRIMARY KEY,
                    task_id TEXT NOT NULL,
                    operator TEXT NOT NULL DEFAULT '',
                    summary TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(task_id) REFERENCES patrol_tasks(task_id) ON DELETE CASCADE
                );

                CREATE TABLE IF NOT EXISTS patrol_task_opinions (
                    opinion_id TEXT PRIMARY KEY,
                    task_id TEXT NOT NULL,
                    project_id TEXT NOT NULL DEFAULT '',
                    stage_id TEXT NOT NULL DEFAULT '',
                    stage_name TEXT NOT NULL DEFAULT '',
                    opinion_no INTEGER NOT NULL DEFAULT 0,
                    title TEXT NOT NULL DEFAULT '',
                    risk_level TEXT NOT NULL DEFAULT '',
                    opinion_content TEXT NOT NULL DEFAULT '',
                    source_files_json TEXT NOT NULL DEFAULT '{}',
                    status TEXT NOT NULL DEFAULT 'pending_photo'
                        CHECK(status IN ('pending_photo','photo_taken','done','returned')),
                    check_opinion TEXT NOT NULL DEFAULT '',
                    return_reason TEXT NOT NULL DEFAULT '',
                    check_by TEXT NOT NULL DEFAULT '',
                    check_time TEXT NOT NULL DEFAULT '',
                    sort INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    FOREIGN KEY(task_id) REFERENCES patrol_tasks(task_id) ON DELETE CASCADE,
                    UNIQUE(task_id, project_id, stage_id, opinion_no)
                );

                CREATE TABLE IF NOT EXISTS patrol_opinion_photos (
                    photo_id TEXT PRIMARY KEY,
                    opinion_id TEXT NOT NULL,
                    file_name TEXT NOT NULL DEFAULT '',
                    file_path TEXT NOT NULL DEFAULT '',
                    taken_at TEXT NOT NULL DEFAULT '',
                    sort INTEGER NOT NULL DEFAULT 0,
                    uploader TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(opinion_id) REFERENCES patrol_task_opinions(opinion_id) ON DELETE CASCADE
                );

                CREATE INDEX IF NOT EXISTS idx_patrol_dict_type
                    ON patrol_dict(dict_type, enabled, sort, created_at);
                CREATE INDEX IF NOT EXISTS idx_patrol_tasks_list
                    ON patrol_tasks(deleted, status, assigned_user_id, created_at DESC);
                CREATE INDEX IF NOT EXISTS idx_patrol_records_task
                    ON patrol_records(task_id, created_at);
                CREATE INDEX IF NOT EXISTS idx_patrol_media_record
                    ON patrol_media(record_id, sort);
                CREATE INDEX IF NOT EXISTS idx_patrol_hazards_task
                    ON patrol_hazards(task_id, status, updated_at DESC);
                CREATE INDEX IF NOT EXISTS idx_patrol_hazard_shots
                    ON patrol_hazard_shots(hazard_id, created_at);
                CREATE INDEX IF NOT EXISTS idx_patrol_task_docs
                    ON patrol_task_docs(task_id, created_at);
                CREATE INDEX IF NOT EXISTS idx_patrol_task_opinions
                    ON patrol_task_opinions(task_id, status, sort, created_at);
                CREATE INDEX IF NOT EXISTS idx_patrol_opinion_photos
                    ON patrol_opinion_photos(opinion_id, sort, created_at);
                CREATE INDEX IF NOT EXISTS idx_patrol_task_lines_task
                    ON patrol_task_lines(task_id, sort);
                CREATE INDEX IF NOT EXISTS idx_patrol_task_users_task
                    ON patrol_task_users(task_id, sort);
                CREATE INDEX IF NOT EXISTS idx_patrol_task_users_uid
                    ON patrol_task_users(user_id);
                """
            )
            # 存量库迁移：为已有 patrol_hazards 补 media_id / video_time 列
            hazard_columns = {
                item["name"]
                for item in connection.execute("PRAGMA table_info(patrol_hazards)").fetchall()
            }
            if "media_id" not in hazard_columns:
                connection.execute("ALTER TABLE patrol_hazards ADD COLUMN media_id TEXT NOT NULL DEFAULT ''")
            if "video_time" not in hazard_columns:
                connection.execute("ALTER TABLE patrol_hazards ADD COLUMN video_time TEXT NOT NULL DEFAULT ''")
            # 存量库迁移：补确认人 / 确认时间列（审计追溯）
            if "confirmer" not in hazard_columns:
                connection.execute("ALTER TABLE patrol_hazards ADD COLUMN confirmer TEXT NOT NULL DEFAULT ''")
            if "confirm_time" not in hazard_columns:
                connection.execute("ALTER TABLE patrol_hazards ADD COLUMN confirm_time TEXT NOT NULL DEFAULT ''")
            # 存量库迁移：补整改提交人 / 提交时间列（审计追溯）
            if "submitter" not in hazard_columns:
                connection.execute("ALTER TABLE patrol_hazards ADD COLUMN submitter TEXT NOT NULL DEFAULT ''")
            if "submit_time" not in hazard_columns:
                connection.execute("ALTER TABLE patrol_hazards ADD COLUMN submit_time TEXT NOT NULL DEFAULT ''")
            # 存量库迁移：为已有 patrol_tasks 补监测方案字段列
            task_columns = {
                item["name"]
                for item in connection.execute("PRAGMA table_info(patrol_tasks)").fetchall()
            }
            # 存量库迁移：为已有 patrol_tasks 补“任务类别”列（默认施工巡查）
            if "task_type" not in task_columns:
                connection.execute("ALTER TABLE patrol_tasks ADD COLUMN task_type TEXT NOT NULL DEFAULT 'construction'")
            for column in (
                "monitor_frequency", "monitor_points", "warning_threshold",
                "emergency_plan", "report_requirement", "review_opinion",
            ):
                if column not in task_columns:
                    connection.execute(f"ALTER TABLE patrol_tasks ADD COLUMN {column} TEXT NOT NULL DEFAULT ''")
            if "source_workflow_id" not in task_columns:
                connection.execute("ALTER TABLE patrol_tasks ADD COLUMN source_workflow_id TEXT NOT NULL DEFAULT ''")
            if "source_project_id" not in task_columns:
                connection.execute("ALTER TABLE patrol_tasks ADD COLUMN source_project_id TEXT NOT NULL DEFAULT ''")
            if "source_project_name" not in task_columns:
                connection.execute("ALTER TABLE patrol_tasks ADD COLUMN source_project_name TEXT NOT NULL DEFAULT ''")
            connection.execute(
                "CREATE UNIQUE INDEX IF NOT EXISTS idx_patrol_tasks_source_workflow "
                "ON patrol_tasks(source_workflow_id) WHERE source_workflow_id <> ''"
            )
            # D122：项目任务唯一索引改为“未删除才占位”，使软删后同项目终审可重建新任务
            connection.execute("DROP INDEX IF EXISTS idx_patrol_tasks_source_project")
            connection.execute(
                "CREATE UNIQUE INDEX IF NOT EXISTS idx_patrol_tasks_source_project "
                "ON patrol_tasks(source_project_id) WHERE source_project_id <> '' AND deleted = 0"
            )
            # 存量库迁移：为已有 patrol_media 补照片名称(caption)列
            media_columns = {
                item["name"]
                for item in connection.execute("PRAGMA table_info(patrol_media)").fetchall()
            }
            if "caption" not in media_columns:
                connection.execute("ALTER TABLE patrol_media ADD COLUMN caption TEXT NOT NULL DEFAULT ''")
            # 存量数据回填：旧单值线路/指派迁移进关联表
            connection.execute(
                "INSERT OR IGNORE INTO patrol_task_lines(task_id, line, sort) "
                "SELECT task_id, line, 0 FROM patrol_tasks WHERE line <> ''"
            )
            connection.execute(
                "INSERT OR IGNORE INTO patrol_task_users(task_id, user_id, user_name, sort) "
                "SELECT task_id, assigned_user_id, assigned_user_name, 0 FROM patrol_tasks "
                "WHERE assigned_user_id <> ''"
            )
            # 基础信息最近修改时间列（仅“编辑/改派”更新；与列表业务活动时间相互独立）
            if "info_updated_at" not in task_columns:
                connection.execute(
                    "ALTER TABLE patrol_tasks ADD COLUMN info_updated_at TEXT NOT NULL DEFAULT ''"
                )
            connection.execute(
                "UPDATE patrol_tasks SET info_updated_at = created_at WHERE info_updated_at = ''"
            )
            # 存量库迁移：为已有 patrol_records 补“本次上传名称”列
            record_columns = {
                item["name"]
                for item in connection.execute("PRAGMA table_info(patrol_records)").fetchall()
            }
            if "title" not in record_columns:
                connection.execute("ALTER TABLE patrol_records ADD COLUMN title TEXT NOT NULL DEFAULT ''")
            # D120：终审综合评价 / 终审人意见（核查栏顶部“综合评价卡”数据源；监测审查意见不再自动预填）
            for d120_column in ("overall_conclusion", "final_opinion"):
                if d120_column not in task_columns:
                    connection.execute(
                        f"ALTER TABLE patrol_tasks ADD COLUMN {d120_column} TEXT NOT NULL DEFAULT ''"
                    )
        self._seed_default_dicts()
        self._migrate_legacy_events()

    def _seed_default_dicts(self) -> None:
        defaults: dict[str, list[str]] = {
            "line": [
                "1号线", "2号线", "3号线", "4号线", "5号线", "6号线", "7号线", "8号线", "9号线", "10号线",
                "S1号线", "S2号线", "S3号线", "S6号线", "S7号线", "S8号线", "S9号线",
            ],
            "construction_type": [
                "基坑开挖", "桩基施工", "盾构穿越", "顶管施工", "降水工程", "爆破作业",
                "堆载", "隧道下穿", "管线迁改", "其他",
            ],
            "hazard_type": ["违规施工", "超范围施工", "未按方案施工", "监测缺失", "防护不足", "其他"],
            "hazard_risk": ["高", "中", "低"],
        }
        with self._lock, self._connect() as connection:
            for dict_type, labels in defaults.items():
                existing = {
                    row["label"]
                    for row in connection.execute(
                        "SELECT label FROM patrol_dict WHERE dict_type = ?", (dict_type,)
                    )
                }
                for index, label in enumerate(labels):
                    if label in existing:
                        continue
                    connection.execute(
                        """
                        INSERT INTO patrol_dict(dict_id, dict_type, label, value, sort, enabled, created_at)
                        VALUES(?, ?, ?, ?, ?, 1, ?)
                        """,
                        (_id("pdict"), dict_type, label, label, index + 1, _now()),
                    )

    def _migrate_legacy_events(self) -> None:
        """把 v1 的 patrol_events/patrol_photos 迁移为 legacy 任务 + 单条巡查记录 + 媒体。"""
        with self._lock, self._connect() as connection:
            tables = {
                row["name"] for row in connection.execute(
                    "SELECT name FROM sqlite_master WHERE type = 'table'"
                )
            }
            if "patrol_events" not in tables:
                return
            rows = connection.execute("SELECT * FROM patrol_events").fetchall()
            for row in rows:
                event = dict(row)
                task_id = _id("ptask")
                now = event["created_at"] or _now()
                connection.execute(
                    """
                    INSERT INTO patrol_tasks(
                        task_id, task_no, name, line, location_desc, requirement,
                        assigned_user_id, assigned_user_name, status, dispatcher, dispatch_time,
                        legacy, deleted, created_at, updated_at
                    ) VALUES(?, ?, ?, ?, ?, '', '', '', 'completed', ?, ?, 1, 0, ?, ?)
                    """,
                    (
                        task_id, event["event_no"], f"历史事件 {event['event_no']}",
                        event["line"], event["location_desc"],
                        event.get("reporter_name") or "", now, now, now,
                    ),
                )
                record_id = _id("prec")
                connection.execute(
                    """
                    INSERT INTO patrol_records(
                        record_id, task_id, type, hazard_id, longitude, latitude, accuracy,
                        note, created_by, created_by_name, created_at
                    ) VALUES(?, ?, 'patrol', '', ?, ?, ?, ?, '', ?, ?)
                    """,
                    (
                        record_id, task_id, event.get("longitude"), event.get("latitude"),
                        event.get("accuracy"), event.get("remark") or "", event.get("reporter_name") or "", now,
                    ),
                )
                photos = connection.execute(
                    "SELECT * FROM patrol_photos WHERE event_id = ? ORDER BY sort, created_at",
                    (event["event_id"],),
                ).fetchall()
                for index, photo in enumerate(photos, start=1):
                    connection.execute(
                        """
                        INSERT INTO patrol_media(
                            media_id, record_id, kind, file_name, file_path, taken_at, sort, created_at
                        ) VALUES(?, ?, 'photo', ?, ?, ?, ?, ?)
                        """,
                        (
                            _id("pmed"), record_id, photo["file_name"], photo["file_path"],
                            photo["taken_at"], index, photo["created_at"] or now,
                        ),
                    )
            connection.execute("ALTER TABLE patrol_events RENAME TO patrol_events_bak")
            connection.execute("ALTER TABLE patrol_photos RENAME TO patrol_photos_bak")

    # ---- 字典 ----

    @staticmethod
    def _dict(row: sqlite3.Row) -> dict[str, Any]:
        return dict(row)

    def list_dicts(self, dict_type: str) -> list[dict[str, Any]]:
        if dict_type not in DICT_TYPES:
            raise ValueError("字典类型无效。")
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT * FROM patrol_dict WHERE dict_type = ? AND enabled = 1 ORDER BY sort, created_at",
                (dict_type,),
            ).fetchall()
        return [self._dict(row) for row in rows]

    def get_dict(self, dict_id: str) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM patrol_dict WHERE dict_id = ?", (dict_id,)).fetchone()
        if row is None:
            raise KeyError("字典项不存在。")
        return self._dict(row)

    def create_dict(self, data: dict[str, Any]) -> dict[str, Any]:
        dict_type = _text(data.get("type"), field="字典类型", required=True, maximum=40)
        if dict_type not in DICT_TYPES:
            raise ValueError("字典类型无效。")
        label = _text(data.get("label"), field="字典名称", required=True, maximum=60)
        value = _text(data.get("value") or label, field="字典值", maximum=80)
        sort = _optional_int(data.get("sort"), field="排序", minimum=0, maximum=10000) or 0
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO patrol_dict(dict_id, dict_type, label, value, sort, enabled, created_at)
                VALUES(?, ?, ?, ?, ?, 1, ?)
                """,
                (_id("pdict"), dict_type, label, value, sort, _now()),
            )
        return self._get_dict_by_label(dict_type, label)

    def _get_dict_by_label(self, dict_type: str, label: str) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM patrol_dict WHERE dict_type = ? AND label = ?", (dict_type, label)
            ).fetchone()
        if row is None:
            raise KeyError("字典项不存在。")
        return self._dict(row)

    def update_dict(self, dict_id: str, data: dict[str, Any]) -> dict[str, Any]:
        current = self.get_dict(dict_id)
        label = _text(data.get("label", current["label"]), field="字典名称", required=True, maximum=60)
        value = _text(data.get("value", current["value"] or current["label"]), field="字典值", maximum=80)
        sort = _optional_int(data.get("sort", current["sort"]), field="排序", minimum=0, maximum=10000)
        enabled = _optional_int(data.get("enabled", current["enabled"]), field="启用", minimum=0, maximum=1)
        with self._lock, self._connect() as connection:
            connection.execute(
                "UPDATE patrol_dict SET label = ?, value = ?, sort = ?, enabled = ? WHERE dict_id = ?",
                (label, value, sort, enabled, dict_id),
            )
        return self.get_dict(dict_id)

    def delete_dict(self, dict_id: str) -> dict[str, Any]:
        current = self.get_dict(dict_id)
        with self._lock, self._connect() as connection:
            connection.execute("DELETE FROM patrol_dict WHERE dict_id = ?", (dict_id,))
        return {"dict_id": dict_id, "label": current["label"], "deleted": True}

    # ---- 任务 ----

    def _next_task_no(self, connection: sqlite3.Connection) -> str:
        prefix = "RW" + _today()
        # 用当日最大序号顺延，而非 COUNT：任务被删除后编号不复用、不冲突
        maximum = connection.execute(
            "SELECT MAX(CAST(SUBSTR(task_no, -4) AS INTEGER)) FROM patrol_tasks WHERE task_no LIKE ?",
            (prefix + "%",),
        ).fetchone()[0]
        return f"{prefix}{int(maximum or 0) + 1:04d}"

    @staticmethod
    def _task(row: sqlite3.Row) -> dict[str, Any]:
        return dict(row)

    def _visible(self, task: dict[str, Any], actor: dict[str, Any]) -> bool:
        if actor.get("is_admin"):
            return True
        assigned = str(task.get("assigned_user_id") or "")
        uid = str(actor.get("user_id") or "")
        if not uid:
            return False
        if assigned == uid:
            return True
        # 多指派：任一关联 user_id 匹配即可见
        with self._connect() as connection:
            return connection.execute(
                "SELECT 1 FROM patrol_task_users WHERE task_id = ? AND user_id = ?",
                (task.get("task_id"), uid),
            ).fetchone() is not None

    def _attach_task_extras(self, tasks: list[dict[str, Any]]) -> None:
        """为任务附加 lines[] / assignees[]，并把旧字段回填为聚合值（兼容展示/水印）。"""
        if not tasks:
            return
        ids = [t["task_id"] for t in tasks]
        with self._connect() as connection:
            ph = ",".join("?" for _ in ids)
            lines: dict[str, list[str]] = {}
            for row in connection.execute(
                f"SELECT task_id, line FROM patrol_task_lines WHERE task_id IN ({ph}) ORDER BY sort, line",
                ids,
            ):
                lines.setdefault(row["task_id"], []).append(row["line"])
            users: dict[str, list[dict[str, str]]] = {}
            for row in connection.execute(
                f"SELECT task_id, user_id, user_name FROM patrol_task_users WHERE task_id IN ({ph}) "
                "ORDER BY sort, user_id",
                ids,
            ):
                users.setdefault(row["task_id"], []).append(
                    {"user_id": row["user_id"], "user_name": row["user_name"]}
                )
        for t in tasks:
            t["lines"] = lines.get(t["task_id"], [])
            t["assignees"] = users.get(t["task_id"], [])
            if t["lines"]:
                t["line"] = "、".join(t["lines"])
            if t["assignees"]:
                t["assigned_user_id"] = ",".join(a["user_id"] for a in t["assignees"])
                t["assigned_user_name"] = "、".join(a["user_name"] or a["user_id"] for a in t["assignees"])

    def _sync_task_assignments(self, task_id: str, lines: list[str], users: list[dict[str, str]]) -> None:
        with self._lock, self._connect() as connection:
            connection.execute("DELETE FROM patrol_task_lines WHERE task_id = ?", (task_id,))
            connection.execute("DELETE FROM patrol_task_users WHERE task_id = ?", (task_id,))
            for i, line in enumerate(lines):
                connection.execute(
                    "INSERT OR IGNORE INTO patrol_task_lines(task_id, line, sort) VALUES(?, ?, ?)",
                    (task_id, line, i),
                )
            for i, u in enumerate(users):
                connection.execute(
                    "INSERT OR IGNORE INTO patrol_task_users(task_id, user_id, user_name, sort) VALUES(?, ?, ?, ?)",
                    (task_id, u.get("user_id", ""), u.get("user_name", ""), i),
                )

    def create_task(
        self,
        data: dict[str, Any],
        actor: dict[str, Any],
        *,
        task_type: str = "construction",
        source_workflow_id: str = "",
        source_project_id: str = "",
    ) -> dict[str, Any]:
        task_id, now = _id("ptask"), _now()
        lines_in = _as_list(data.get("lines") or data.get("line_list"))
        uids_in = _as_list(data.get("user_ids") or data.get("assigned_user_ids"))
        names_in = _as_list(data.get("user_names") or data.get("assigned_user_names"))
        name = _text(data.get("name"), field="任务名称", required=True, maximum=120)
        line = _text(data.get("line"), field="线路", maximum=60)
        location_desc = _text(data.get("location_desc"), field="位置描述", maximum=200)
        requirement = _text(data.get("requirement"), field="巡查内容", maximum=2000)
        assigned_user_id = _text(data.get("assigned_user_id"), field="指派账号", maximum=200)
        assigned_user_name = _text(data.get("assigned_user_name"), field="指派账号名称", maximum=200)
        # 旧列存首值便于兼容（展示统一走 lines/assignees 关联值）
        if not line and lines_in:
            line = lines_in[0]
        if not assigned_user_id and uids_in:
            assigned_user_id = uids_in[0]
            assigned_user_name = names_in[0] if names_in else uids_in[0]
        remark = _text(data.get("remark"), field="备注", maximum=1000)
        monitor_frequency = _text(data.get("monitor_frequency"), field="监测频率", maximum=2000)
        monitor_points = _text(data.get("monitor_points"), field="监测点位", maximum=2000)
        warning_threshold = _text(data.get("warning_threshold"), field="预警阈值", maximum=2000)
        emergency_plan = _text(data.get("emergency_plan"), field="应急预案", maximum=2000)
        report_requirement = _text(data.get("report_requirement"), field="数据报送要求", maximum=2000)
        review_opinion = _text(data.get("review_opinion"), field="监测审查意见", maximum=2000)
        dispatcher = _text(actor.get("name") or actor.get("user_id"), field="派发人", maximum=60)
        source_workflow_id = _text(source_workflow_id, field="审核流程", maximum=80)
        source_project_id = _text(source_project_id, field="项目档案", maximum=80)
        source_project_name = _text(data.get("source_project_name") or name, field="项目名称", maximum=120)
        task_type = _text(task_type, field="任务类别", maximum=20)
        if task_type not in {"pre_construction", "construction"}:
            raise ValueError("任务类别无效。")
        with self._lock, self._connect() as connection:
            task_no = self._next_task_no(connection)
            connection.execute(
                """
                INSERT INTO patrol_tasks(
                    task_id, task_no, task_type, source_workflow_id, source_project_id, source_project_name,
                    name, line, location_desc, requirement,
                    assigned_user_id, assigned_user_name, status, dispatcher, dispatch_time,
                    remark, monitor_frequency, monitor_points, warning_threshold,
                    emergency_plan, report_requirement, review_opinion,
                    legacy, deleted, created_at, updated_at
                ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, 0, ?, ?)
                """,
                (
                    task_id, task_no, task_type, source_workflow_id, source_project_id, source_project_name,
                    name, line, location_desc, requirement,
                    assigned_user_id, assigned_user_name, dispatcher, now, remark,
                    monitor_frequency, monitor_points, warning_threshold,
                    emergency_plan, report_requirement, review_opinion,
                    now, now,
                ),
            )
        # 多线路 / 多巡查员关联行
        final_lines = lines_in or ([line] if line else [])
        final_uids = uids_in or ([assigned_user_id] if assigned_user_id else [])
        final_names = names_in if names_in else ([assigned_user_name] if assigned_user_name else [])
        users = [
            {"user_id": u, "user_name": (final_names[i] if i < len(final_names) else u)}
            for i, u in enumerate(final_uids)
        ]
        self._sync_task_assignments(task_id, final_lines, users)
        return self.get_task(task_id, actor)

    def create_workflow_task(self, data: dict[str, Any]) -> dict[str, Any]:
        """终审通过后建/复用“施工前核查”任务，并把本次终审意见结构化写入“审核意见现场核查”。

        D120：逐条意见一律落 patrol_task_opinions（含未绑定项目档案的流程）；巡查内容只放
        “指引+条数”；监测审查意见(review_opinion)不再自动预填长文；综合评价/终审人意见写入
        overall_conclusion/final_opinion 供详情页“综合评价卡”展示。
        幂等：有 project_id 按项目唯一，否则按流程(source_workflow_id)唯一。
        """
        system_actor = {"user_id": "workflow", "name": "审核流转系统", "is_admin": True}
        project_id = _text(data.get("project_id"), field="项目档案", maximum=80)
        workflow_items = self._workflow_opinion_items(data, project_id=project_id)
        # D122：终审通过创建时可提交可选参数（巡查员多选/位置/备注）。未带 user_ids 键（非弹窗调用）
        # 时回退为指派流程发起人，保持向后兼容；显式传空数组 = 不指派任何人。
        location_desc = _text(data.get("location_desc"), field="位置描述", maximum=200)
        user_remark = _text(data.get("remark"), field="备注", maximum=1000)
        has_user_key = any(key in data for key in ("user_ids", "assigned_user_ids"))
        if has_user_key:
            uid_values = _as_list(data.get("user_ids") or data.get("assigned_user_ids"))
            name_values = _as_list(data.get("user_names") or data.get("assigned_user_names"))
            if not name_values:
                name_values = [str(value) for value in uid_values]
            name_values = name_values[: len(uid_values)]
        else:
            fallback_uid = _text(data.get("assigned_user_id") or data.get("initiator_id"), field="指派账号", maximum=200)
            fallback_name = _text(data.get("assigned_user_name") or data.get("initiator_name"), field="指派账号名称", maximum=200)
            uid_values = [fallback_uid] if fallback_uid else []
            name_values = [fallback_name or fallback_uid] if fallback_uid else []
        task_id = ""
        created = False
        if project_id:
            task, created = self.ensure_project_task(
                project_id,
                project_name=data.get("project_name") or data.get("projectName") or "",
                workflow_id=data.get("workflow_id") or "",
                initiator_id=data.get("initiator_id"),
                initiator_name=data.get("initiator_name"),
            )
            task_id = task["task_id"]
        else:
            workflow_id = _text(data.get("workflow_id"), field="审核流程", required=True, maximum=80)
            with self._connect() as connection:
                existing = connection.execute(
                    "SELECT task_id FROM patrol_tasks WHERE source_workflow_id = ? AND deleted = 0",
                    (workflow_id,),
                ).fetchone()
            if existing is not None:
                task_id = existing["task_id"]
            else:
                project_name = _text(data.get("project_name"), field="项目名称", required=True, maximum=120)
                default_remark = f"来源：案例审核流转终审通过（流程 {workflow_id}）。"
                task_data = {
                    "name": project_name,
                    "requirement": _patrol_requirement_text(len(workflow_items)),
                    "user_ids": uid_values,
                    "user_names": name_values,
                    "location_desc": location_desc,
                    "remark": _limit_text(user_remark or default_remark, 1000),
                }
                try:
                    task = self.create_task(
                        task_data, system_actor, task_type="pre_construction",
                        source_workflow_id=workflow_id,
                    )
                    task_id = task["task_id"]
                    created = True
                except sqlite3.IntegrityError:
                    with self._connect() as connection:
                        existing = connection.execute(
                            "SELECT task_id FROM patrol_tasks WHERE source_workflow_id = ? AND deleted = 0",
                            (workflow_id,),
                        ).fetchone()
                    if existing is None:
                        raise
                    task_id = existing["task_id"]
        if not task_id:
            raise ValueError("巡查任务创建失败。")
        # D122：新建（项目路径）时应用终审弹窗提交的创建参数
        if created and project_id:
            self._apply_creation_persona(
                task_id,
                override_assign=has_user_key,
                uids=uid_values,
                names=name_values,
                location=location_desc,
                remark=user_remark,
            )
        # 结构化意见：项目路径 = 档案历史累计 + 本次；流程路径 = 仅本次（同样落核查栏）
        added = 0
        warning = ""
        try:
            added = self.sync_project_opinions(
                task_id,
                project_id if project_id else "",
                workflow_items=workflow_items,
            )
        except Exception:
            LOGGER.exception("终审后同步意见失败 task=%s project=%s", task_id, project_id)
            warning = "任务已创建，但本次意见同步失败，请在任务详情点「同步」重试。"
        # 终审元数据：综合评价卡 + 巡查内容形态 + 监测审查意见不预填（旧长文存量自动迁移）
        self._finalize_workflow_task(task_id, _workflow_overall_meta(data), created)
        result = {
            "created": created,
            "opinions_added": added,
            "task": self.get_task(task_id, system_actor),
        }
        if warning:
            result["warning"] = warning
        return result

    def _finalize_workflow_task(
        self,
        task_id: str,
        meta: dict[str, str],
        created: bool,
    ) -> None:
        """写入综合评价/终审人意见；巡查内容收敛为“指引+条数”；监测审查意见不自动预填。"""
        with self._lock, self._connect() as connection:
            row = connection.execute(
                "SELECT requirement, review_opinion FROM patrol_tasks WHERE task_id = ?",
                (task_id,),
            ).fetchone()
            if row is None:
                return
            current_requirement = row["requirement"] or ""
            current_review = row["review_opinion"] or ""

            def _legacy_blob(value: str) -> bool:
                # 识别旧逻辑整段拼接终审长文的形态
                return bool(value) and (
                    value.startswith("请依据以下终审意见开展现场符合性巡查")
                    or value.startswith("终审综合结论：")
                    or "；1. " in value
                    or "；2. " in value
                )

            count = connection.execute(
                "SELECT COUNT(*) AS c FROM patrol_task_opinions WHERE task_id = ?", (task_id,)
            ).fetchone()["c"]
            template = _patrol_requirement_text(count)
            default_project_requirement = (
                "请按项目档案中记录的审核意见逐项开展现场符合性巡查：实地核实每条意见所涉情况，"
                "拍摄现场照片并对照说明。"
            )
            new_requirement = current_requirement
            if created or _legacy_blob(current_requirement) or current_requirement == default_project_requirement or not current_requirement.strip():
                new_requirement = template
            new_review = current_review
            if created or _legacy_blob(current_review):
                new_review = ""
            connection.execute(
                """
                UPDATE patrol_tasks SET requirement = ?, review_opinion = ?,
                    overall_conclusion = ?, final_opinion = ?, updated_at = ?
                WHERE task_id = ?
                """,
                (
                    new_requirement, new_review,
                    meta.get("overall_conclusion", ""), meta.get("final_opinion", ""),
                    _now(), task_id,
                ),
            )

    # ---- 审核意见 → 现场核查（Step①）----

    def ensure_project_task(
        self,
        project_id: str,
        *,
        project_name: str = "",
        workflow_id: str = "",
        initiator_id: Any = "",
        initiator_name: Any = "",
    ) -> tuple[dict[str, Any], bool]:
        """按项目找/建唯一巡查任务（source_project_id 唯一索引）。返回 (task, created)。"""
        system_actor = {"user_id": "workflow", "name": "审核流转系统", "is_admin": True}
        with self._connect() as connection:
            existing = connection.execute(
                "SELECT task_id FROM patrol_tasks WHERE source_project_id = ? AND deleted = 0",
                (project_id,),
            ).fetchone()
        if existing is not None:
            return self.get_task(existing["task_id"], system_actor), False
        name = _text(project_name, field="项目名称", required=True, maximum=120)
        task_data = {
            "name": name,
            "source_project_name": name,
            "requirement": "请按项目档案中记录的审核意见逐项开展现场符合性巡查：实地核实每条意见所涉情况，拍摄现场照片并对照说明。",
            "assigned_user_id": str(initiator_id or ""),
            "assigned_user_name": str(initiator_name or ""),
            "remark": f"来源：项目档案（{project_id}）。",
        }
        try:
            task = self.create_task(
                task_data, system_actor, task_type="pre_construction",
                source_workflow_id=_text(workflow_id, field="审核流程", maximum=80),
                source_project_id=project_id,
            )
            return task, True
        except sqlite3.IntegrityError:
            # 并发兜底：项目任务已被其它调用建立
            with self._connect() as connection:
                existing = connection.execute(
                    "SELECT task_id FROM patrol_tasks WHERE source_project_id = ? AND deleted = 0",
                    (project_id,),
                ).fetchone()
            if existing is None:
                raise
            return self.get_task(existing["task_id"], system_actor), False

    def _apply_creation_persona(
        self,
        task_id: str,
        *,
        override_assign: bool = False,
        uids: list[str] | None = None,
        names: list[str] | None = None,
        location: str = "",
        remark: str = "",
    ) -> None:
        """把终审通过时提交的可选创建参数应用到新建任务（仅新建调用）。"""
        uid_values = [str(value) for value in (uids or [])]
        name_values = [str(value) for value in (names or [])]
        if uid_values and not name_values:
            name_values = [value for value in uid_values]
        name_values = name_values[: len(uid_values)]
        assigned_id = uid_values[0] if uid_values else ""
        assigned_name = name_values[0] if name_values else assigned_id
        sets: list[str] = []
        params: list[Any] = []
        if override_assign:
            sets.extend(["assigned_user_id = ?", "assigned_user_name = ?"])
            params.extend([assigned_id, assigned_name])
        if location:
            sets.append("location_desc = ?")
            params.append(_limit_text(location, 200))
        if remark:
            sets.append("remark = ?")
            params.append(_limit_text(remark, 1000))
        if not sets:
            return
        params.append(_now())
        params.append(task_id)
        with self._lock, self._connect() as connection:
            connection.execute(
                f"UPDATE patrol_tasks SET {', '.join(sets)}, updated_at = ? WHERE task_id = ?",
                params,
            )
        if override_assign:
            with self._connect() as connection:
                line_rows = connection.execute(
                    "SELECT line FROM patrol_task_lines WHERE task_id = ? ORDER BY sort, line",
                    (task_id,),
                ).fetchall()
            lines = [row["line"] for row in line_rows]
            users = [
                {"user_id": value, "user_name": name_values[i] if i < len(name_values) else value}
                for i, value in enumerate(uid_values)
            ]
            self._sync_task_assignments(task_id, lines, users)

    def get_project_task(self, project_id: str, actor: dict[str, Any]) -> dict[str, Any] | None:
        """按项目查唯一关联任务（供档案页跳转；软删的不算）。"""
        system_actor = {"user_id": "workflow", "name": "审核流转系统", "is_admin": True}
        with self._connect() as connection:
            row = connection.execute(
                "SELECT task_id FROM patrol_tasks WHERE source_project_id = ? AND deleted = 0 "
                "ORDER BY created_at LIMIT 1",
                (project_id,),
            ).fetchone()
        if row is None:
            return None
        task = self.get_task(row["task_id"], system_actor)
        if not actor.get("is_admin") and not self._visible(task, actor):
            return None
        return task

    def _workflow_opinion_items(self, data: dict[str, Any], project_id: str = "") -> list[dict[str, Any]]:
        """从终审通过请求里解析“本次流程”的逐条审核意见（档案尚未归档时也可下发）。

        stage 信息以项目档案真实阶段为准（防流程快照名称如 AI 猜测的“设计”污染来源 tag）。
        """
        raw = data.get("latest_result_json")
        if isinstance(raw, str) and raw.strip():
            try:
                raw = json.loads(raw)
            except json.JSONDecodeError:
                raw = {}
        payload = raw if isinstance(raw, dict) else {}
        stage_id = _text(data.get("stage_id") or payload.get("stage_id"), field="阶段", maximum=80)
        stage_name = _text(data.get("stage_name") or payload.get("stage_name"), field="阶段名称", maximum=80)
        stage_id, stage_name = _resolve_archive_stage(project_id, stage_id, stage_name)
        return [
            {
                "stage_id": stage_id,
                "stage_name": stage_name,
                "title": title,
                "risk_level": risk,
                "content": content,
            }
            for title, risk, content in _opinion_result_items(payload)
        ]

    def sync_project_opinions(
        self,
        task_id: str,
        project_id: str,
        workflow_items: list[dict[str, Any]] | None = None,
    ) -> int:
        """把项目档案中该项目已审核记录的意见 + 本次终审通过意见幂等合并进任务（只新增不改旧）。

        返回新增条数。意见序号按阶段内顺序连续编号；同阶段同内容自动去重。
        """
        rows = _collect_project_opinions(project_id)
        rows.extend(workflow_items or [])
        if not rows:
            return 0
        now = _now()
        added = 0
        seen: set[str] = set()
        # 先去掉同阶段完全重复的意见（如“终审本次意见”与“档案同阶段已归档记录”重复）
        unique_rows: list[dict[str, Any]] = []
        for item in rows:
            key = "{}|{}|{}|{}".format(
                item.get("stage_id", ""),
                item.get("title", ""),
                item.get("content", ""),
                item.get("risk_level", ""),
            )[:500]
            if key in seen:
                continue
            seen.add(key)
            unique_rows.append(item)
        # 阶段内连续编号（幂等：UNIQUE(task_id, project_id, stage_id, opinion_no)）
        stage_counters: dict[str, int] = {}
        prepared: list[tuple[Any, ...]] = []
        for item in unique_rows:
            stage_id = item.get("stage_id", "") or ""
            stage_counters[stage_id] = stage_counters.get(stage_id, 0) + 1
            prepared.append((item, stage_counters[stage_id]))
        with self._lock, self._connect() as connection:
            for item, opinion_no in prepared:
                opinion_id = _id("popi")
                inserted = connection.execute(
                    """
                    INSERT OR IGNORE INTO patrol_task_opinions(
                        opinion_id, task_id, project_id, stage_id, stage_name,
                        opinion_no, title, risk_level, opinion_content, sort,
                        status, created_at, updated_at
                    ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending_photo', ?, ?)
                    """,
                    (
                        opinion_id, task_id, project_id,
                        item.get("stage_id", ""), item.get("stage_name", ""),
                        opinion_no, item.get("title", ""),
                        item.get("risk_level", ""), item.get("content", ""),
                        opinion_no, now, now,
                    ),
                )
                added += 1 if inserted.rowcount else 0
        return added

    def get_opinion(self, opinion_id: str) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM patrol_task_opinions WHERE opinion_id = ?", (opinion_id,)
            ).fetchone()
        if row is None:
            raise KeyError("核查意见不存在。")
        return dict(row)

    def _load_opinion_photos(self, opinion_ids: list[str]) -> dict[str, list[dict[str, Any]]]:
        photo_map: dict[str, list[dict[str, Any]]] = {}
        if not opinion_ids:
            return photo_map
        with self._connect() as connection:
            placeholders = ",".join("?" for _ in opinion_ids)
            rows = connection.execute(
                f"SELECT * FROM patrol_opinion_photos WHERE opinion_id IN ({placeholders}) "
                "ORDER BY sort, created_at",
                opinion_ids,
            ).fetchall()
        for row in rows:
            photo_map.setdefault(row["opinion_id"], []).append(dict(row))
        return photo_map

    def list_task_opinions(self, task_id: str, actor: dict[str, Any]) -> list[dict[str, Any]]:
        """列出任务下的核查意见（含照片），并对项目类任务做一次幂等同步。"""
        task = self.get_task(task_id, actor)
        project_id = task.get("source_project_id") or ""
        if project_id:
            self.sync_project_opinions(task_id, project_id)
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT * FROM patrol_task_opinions WHERE task_id = ? "
                "ORDER BY sort, created_at, opinion_id",
                (task_id,),
            ).fetchall()
        opinions = [dict(row) for row in rows]
        photo_map = self._load_opinion_photos([o["opinion_id"] for o in opinions])
        for opinion in opinions:
            opinion["photos"] = photo_map.get(opinion["opinion_id"], [])
        return opinions

    def opinion_photo_path(self, photo_id: str) -> Path:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM patrol_opinion_photos WHERE photo_id = ?", (photo_id,)
            ).fetchone()
        if row is None:
            raise KeyError("意见照片不存在。")
        path = Path(row["file_path"]).resolve()
        if not path.exists() or not path.is_file():
            raise KeyError("意见照片文件缺失。")
        return path

    def add_opinion_photo(
        self, opinion_id: str, stored_path: str | Path, file_name: str, taken_at: str, actor: dict[str, Any]
    ) -> dict[str, Any]:
        opinion = self.get_opinion(opinion_id)
        self.get_task(opinion["task_id"], actor)  # 行级隔离
        if opinion["status"] == "done":
            raise ValueError("该意见已完成，照片已锁定，不能再上传。")
        now = _now()
        with self._lock, self._connect() as connection:
            count = int(connection.execute(
                "SELECT COUNT(*) FROM patrol_opinion_photos WHERE opinion_id = ?", (opinion_id,)
            ).fetchone()[0])
            if count >= MAX_OPINION_PHOTOS:
                raise ValueError(f"每条意见最多{MAX_OPINION_PHOTOS}张现场照片。")
            photo_id = _id("pphoto")
            connection.execute(
                """
                INSERT INTO patrol_opinion_photos(
                    photo_id, opinion_id, file_name, file_path, taken_at, sort, uploader, created_at
                ) VALUES(?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (photo_id, opinion_id, file_name, str(stored_path), taken_at, count + 1,
                 actor.get("name") or actor.get("user_id") or "", now),
            )
            connection.execute(
                "UPDATE patrol_task_opinions SET updated_at = ? WHERE opinion_id = ?", (now, opinion_id)
            )
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM patrol_opinion_photos WHERE photo_id = ?", (photo_id,)
            ).fetchone()
        return dict(row)

    def delete_opinion_photo(self, photo_id: str, actor: dict[str, Any]) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM patrol_opinion_photos WHERE photo_id = ?", (photo_id,)
            ).fetchone()
        if row is None:
            raise KeyError("意见照片不存在。")
        opinion = self.get_opinion(row["opinion_id"])
        task = self.get_task(opinion["task_id"], actor)
        if opinion["status"] == "done":
            raise ValueError("该意见已完成，照片已锁定，不能删除。")
        if not actor.get("is_admin") and str(actor.get("user_id") or "") != str(task.get("assigned_user_id") or ""):
            raise ValueError("无权限删除该照片。")
        with self._lock, self._connect() as connection:
            connection.execute("DELETE FROM patrol_opinion_photos WHERE photo_id = ?", (photo_id,))
            connection.execute(
                "UPDATE patrol_task_opinions SET updated_at = ? WHERE opinion_id = ?",
                (_now(), opinion["opinion_id"]),
            )
        Path(row["file_path"]).unlink(missing_ok=True)
        return {"photo_id": photo_id, "deleted": True}

    def submit_opinion(self, opinion_id: str, actor: dict[str, Any]) -> dict[str, Any]:
        opinion = self.get_opinion(opinion_id)
        self.get_task(opinion["task_id"], actor)
        if opinion["status"] == "done":
            raise ValueError("该意见已完成，无需再提交。")
        with self._connect() as connection:
            count = int(connection.execute(
                "SELECT COUNT(*) FROM patrol_opinion_photos WHERE opinion_id = ?", (opinion_id,)
            ).fetchone()[0])
        if count <= 0:
            raise ValueError("请先上传现场照片再提交。")
        with self._lock, self._connect() as connection:
            connection.execute(
                "UPDATE patrol_task_opinions SET status = 'photo_taken', updated_at = ? "
                "WHERE opinion_id = ?",
                (_now(), opinion_id),
            )
        return self.get_opinion(opinion_id)

    def update_opinion_check(self, opinion_id: str, comment: str, actor: dict[str, Any]) -> dict[str, Any]:
        if not actor.get("is_admin"):
            raise ValueError("无权限编辑核查意见。")
        opinion = self.get_opinion(opinion_id)
        self.get_task(opinion["task_id"], actor)
        comment_value = _text(comment, field="核查意见", maximum=2000)
        with self._lock, self._connect() as connection:
            connection.execute(
                "UPDATE patrol_task_opinions SET check_opinion = ?, updated_at = ? WHERE opinion_id = ?",
                (comment_value, _now(), opinion_id),
            )
        return self.get_opinion(opinion_id)

    def review_opinion(self, opinion_id: str, data: dict[str, Any], actor: dict[str, Any]) -> dict[str, Any]:
        if not actor.get("is_admin"):
            raise ValueError("无权限复核现场核查结果。")
        opinion = self.get_opinion(opinion_id)
        self.get_task(opinion["task_id"], actor)
        result = _text(data.get("result"), field="核查结论", required=True, maximum=20)
        if result not in {"pass", "return"}:
            raise ValueError("核查结论无效。")
        comment = _text(data.get("comment"), field="核查说明", maximum=2000)
        now = _now()
        name = actor.get("name") or actor.get("user_id") or ""
        if result == "pass":
            new_status, return_reason = "done", ""
        else:
            new_status, return_reason = "returned", comment
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                UPDATE patrol_task_opinions SET status = ?, return_reason = ?,
                    check_opinion = ?, check_by = ?, check_time = ?, updated_at = ?
                WHERE opinion_id = ?
                """,
                (new_status, return_reason, comment, name, now, now, opinion_id),
            )
        return self.get_opinion(opinion_id)

    def get_task(self, task_id: str, actor: dict[str, Any]) -> dict[str, Any]:
        with self._connect() as connection:
            # 已软删除的任务对所有人不可见，阻断详情读取与后续全部写入路径
            row = connection.execute(
                "SELECT * FROM patrol_tasks WHERE task_id = ? AND deleted = 0", (task_id,)
            ).fetchone()
            if row is None:
                raise KeyError("巡查任务不存在。")
            records = connection.execute(
                "SELECT * FROM patrol_records WHERE task_id = ? ORDER BY created_at, record_id", (task_id,)
            ).fetchall()
            hazards = connection.execute(
                "SELECT * FROM patrol_hazards WHERE task_id = ? ORDER BY created_at, hazard_id", (task_id,)
            ).fetchall()
            record_ids = [record["record_id"] for record in records]
        task = self._task(row)
        if not self._visible(task, actor):
            raise KeyError("巡查任务不存在。")
        media_map: dict[str, list[dict[str, Any]]] = {}
        if record_ids:
            with self._connect() as connection:
                placeholders = ",".join("?" for _ in record_ids)
                media_rows = connection.execute(
                    f"SELECT * FROM patrol_media WHERE record_id IN ({placeholders}) ORDER BY sort, created_at",
                    record_ids,
                ).fetchall()
            for media in media_rows:
                media_map.setdefault(media["record_id"], []).append(dict(media))
        task["records"] = [
            {**dict(record), "media": media_map.get(record["record_id"], [])} for record in records
        ]
        shot_map: dict[str, list[dict[str, Any]]] = {}
        hazard_ids = [hazard["hazard_id"] for hazard in hazards]
        if hazard_ids:
            with self._connect() as connection:
                placeholders = ",".join("?" for _ in hazard_ids)
                shot_rows = connection.execute(
                    f"SELECT * FROM patrol_hazard_shots WHERE hazard_id IN ({placeholders}) ORDER BY created_at, shot_id",
                    hazard_ids,
                ).fetchall()
            for shot in shot_rows:
                shot_map.setdefault(shot["hazard_id"], []).append(dict(shot))
        task["hazards"] = [
            {**dict(hazard), "shots": shot_map.get(hazard["hazard_id"], [])} for hazard in hazards
        ]
        with self._connect() as connection:
            doc_rows = connection.execute(
                "SELECT * FROM patrol_task_docs WHERE task_id = ? ORDER BY created_at, doc_id", (task_id,)
            ).fetchall()
        task["docs"] = [dict(doc) for doc in doc_rows]
        with self._connect() as connection:
            opinion_rows = connection.execute(
                "SELECT * FROM patrol_task_opinions WHERE task_id = ? "
                "ORDER BY sort, created_at, opinion_id",
                (task_id,),
            ).fetchall()
        opinions = [dict(opinion) for opinion in opinion_rows]
        if opinions:
            photo_map = self._load_opinion_photos([o["opinion_id"] for o in opinions])
            for opinion in opinions:
                opinion["photos"] = photo_map.get(opinion["opinion_id"], [])
        task["opinions"] = opinions
        with self._connect() as connection:
            log_rows = connection.execute(
                "SELECT * FROM patrol_task_edit_logs WHERE task_id = ? "
                "ORDER BY created_at DESC, log_id LIMIT 30", (task_id,)
            ).fetchall()
        task["edit_logs"] = [dict(log) for log in log_rows]
        self._attach_task_extras([task])
        return task

    def list_tasks(
        self,
        actor: dict[str, Any],
        *,
        page: int = 1,
        size: int = 20,
        line: str = "",
        status_value: str = "",
        task_type_value: str = "",
        assigned_user_id: str = "",
        date_from: str = "",
        date_to: str = "",
        keyword: str = "",
        order_by: str = "created",
    ) -> dict[str, Any]:
        page = max(1, int(page))
        size = max(1, min(100, int(size)))
        clauses = ["deleted = 0"]
        params: list[Any] = []
        if not actor.get("is_admin"):
            user_id = str(actor.get("user_id") or "")
            if not user_id:
                # 未提供有效身份的请求不应看到任何任务（含历史空指派任务）
                clauses.append("1 = 0")
            else:
                clauses.append(
                    "(EXISTS (SELECT 1 FROM patrol_task_users tu WHERE tu.task_id = {ref}.task_id AND tu.user_id = ?) "
                    "OR assigned_user_id = ?)"
                )
                params.extend([user_id, user_id])
        elif assigned_user_id.strip():
            clauses.append(
                "EXISTS (SELECT 1 FROM patrol_task_users tu WHERE tu.task_id = {ref}.task_id AND tu.user_id = ?)"
            )
            params.append(assigned_user_id.strip())
        if line.strip():
            clauses.append(
                "EXISTS (SELECT 1 FROM patrol_task_lines pl WHERE pl.task_id = {ref}.task_id AND pl.line = ?)"
            )
            params.append(line.strip())
        if status_value in TASK_STATUSES:
            clauses.append("status = ?")
            params.append(status_value)
        if task_type_value in {"pre_construction", "construction"}:
            clauses.append("task_type = ?")
            params.append(task_type_value)
        if keyword.strip():
            clauses.append("(name LIKE ? OR task_no LIKE ?)")
            params.extend([f"%{keyword.strip()}%", f"%{keyword.strip()}%"])
        if date_from.strip():
            clauses.append("substr(created_at, 1, 10) >= ?")
            params.append(date_from.strip())
        if date_to.strip():
            clauses.append("substr(created_at, 1, 10) <= ?")
            params.append(date_to.strip())
        where_count = "WHERE " + " AND ".join(c.replace("{ref}", "patrol_tasks") for c in clauses)
        where_list = "WHERE " + " AND ".join(c.replace("{ref}", "t") for c in clauses)
        recent_expr = ("MAX(t.created_at, t.updated_at, "
                       "COALESCE((SELECT MAX(r.created_at) FROM patrol_records r WHERE r.task_id = t.task_id), ''), "
                       "COALESCE((SELECT MAX(h.updated_at) FROM patrol_hazards h WHERE h.task_id = t.task_id), ''), "
                       "COALESCE((SELECT MAX(o.updated_at) FROM patrol_task_opinions o WHERE o.task_id = t.task_id), ''))")
        order_sql = (recent_expr + " DESC, t.task_id") if order_by == "recent" else "t.created_at DESC, t.task_id"
        with self._connect() as connection:
            total = int(connection.execute(
                f"SELECT COUNT(*) FROM patrol_tasks {where_count}", params
            ).fetchone()[0])
            rows = connection.execute(
                f"""
                SELECT t.*,
                       (SELECT COUNT(*) FROM patrol_hazards h WHERE h.task_id = t.task_id) AS hazard_count,
                       (SELECT COUNT(*) FROM patrol_hazards h WHERE h.task_id = t.task_id AND h.status != 'closed') AS open_hazard_count,
                       (SELECT COUNT(*) FROM patrol_records r WHERE r.task_id = t.task_id) AS record_count,
                       {recent_expr} AS recent_time
                FROM patrol_tasks t
                {where_list}
                ORDER BY {order_sql}
                LIMIT ? OFFSET ?
                """,
                [*params, size, (page - 1) * size],
            ).fetchall()
        items = []
        for row in rows:
            task = self._task(row)
            task["hazard_count"] = row["hazard_count"]
            task["record_count"] = row["record_count"]
            task["recent_time"] = row["recent_time"] or task["created_at"]
            items.append(task)
        self._attach_task_extras(items)
        return {"total": total, "page": page, "size": size, "items": items}

    def update_task(self, task_id: str, data: dict[str, Any], actor: dict[str, Any]) -> dict[str, Any]:
        current = self.get_task(task_id, actor)
        if not actor.get("is_admin"):
            raise ValueError("无权限更新任务。")
        # 多值：优先数组（前端编辑必带），否则沿用现状
        has_lines = data.get("lines") is not None or data.get("line_list") is not None
        has_users = (data.get("user_ids") is not None or data.get("assigned_user_ids") is not None)
        lines_in = _as_list(data.get("lines") or data.get("line_list")) if has_lines else list(current.get("lines") or [])
        uids_in = _as_list(data.get("user_ids") or data.get("assigned_user_ids")) if has_users else [
            a["user_id"] for a in (current.get("assignees") or [])
        ]
        names_in = _as_list(data.get("user_names") or data.get("assigned_user_names")) if has_users else [
            a.get("user_name", "") for a in (current.get("assignees") or [])
        ]
        name = _text(data.get("name", current["name"]), field="任务名称", required=True, maximum=120)
        line = _text(data.get("line", current["line"]), field="线路", maximum=200)
        location_desc = _text(data.get("location_desc", current["location_desc"]), field="位置描述", maximum=200)
        requirement = _text(data.get("requirement", current["requirement"]), field="巡查内容", maximum=2000)
        assigned_user_id = _text(data.get("assigned_user_id", current["assigned_user_id"]), field="指派账号", maximum=200)
        assigned_user_name = _text(data.get("assigned_user_name", current["assigned_user_name"]), field="指派账号名称", maximum=200)
        if not line and lines_in:
            line = lines_in[0]
        if not assigned_user_id and uids_in:
            assigned_user_id = uids_in[0]
            assigned_user_name = names_in[0] if names_in else uids_in[0]
        monitor_frequency = _text(data.get("monitor_frequency", current["monitor_frequency"]), field="监测频率", maximum=2000)
        monitor_points = _text(data.get("monitor_points", current["monitor_points"]), field="监测点位", maximum=2000)
        warning_threshold = _text(data.get("warning_threshold", current["warning_threshold"]), field="预警阈值", maximum=2000)
        emergency_plan = _text(data.get("emergency_plan", current["emergency_plan"]), field="应急预案", maximum=2000)
        report_requirement = _text(data.get("report_requirement", current["report_requirement"]), field="数据报送要求", maximum=2000)
        review_opinion = _text(data.get("review_opinion", current["review_opinion"]), field="监测审查意见", maximum=2000)
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                UPDATE patrol_tasks SET name = ?, line = ?, location_desc = ?, requirement = ?,
                    assigned_user_id = ?, assigned_user_name = ?,
                    monitor_frequency = ?, monitor_points = ?, warning_threshold = ?,
                    emergency_plan = ?, report_requirement = ?, review_opinion = ?,
                    updated_at = ?, info_updated_at = ?
                WHERE task_id = ?
                """,
                (name, line, location_desc, requirement, assigned_user_id, assigned_user_name,
                 monitor_frequency, monitor_points, warning_threshold,
                 emergency_plan, report_requirement, review_opinion, _now(), _now(), task_id),
            )
        final_lines = lines_in or ([line] if line else [])
        final_uids = uids_in or ([assigned_user_id] if assigned_user_id else [])
        final_names = names_in if names_in else ([assigned_user_name] if assigned_user_name else [])
        users = [
            {"user_id": u, "user_name": (final_names[i] if i < len(final_names) else u)}
            for i, u in enumerate(final_uids)
        ]
        self._sync_task_assignments(task_id, final_lines, users)
        # 基础信息编辑留痕（记录修改字段）
        info_keys = [k for k in ("name", "lines", "user_ids", "location_desc", "requirement")
                     if data.get(k) is not None]
        if info_keys:
            now_log = _now()
            summary = "，".join(_FIELD_CN.get(k, k) for k in info_keys)
            with self._lock, self._connect() as connection:
                connection.execute(
                    "INSERT INTO patrol_task_edit_logs(log_id, task_id, operator, summary, created_at) "
                    "VALUES(?, ?, ?, ?, ?)",
                    (_id("pedit"), task_id, actor.get("name") or actor.get("user_id") or "", summary, now_log),
                )
        return self.get_task(task_id, actor)

    def set_task_status(self, task_id: str, status_value: str, actor: dict[str, Any]) -> dict[str, Any]:
        current = self.get_task(task_id, actor)
        if not actor.get("is_admin"):
            raise ValueError("无权限变更任务状态。")
        if status_value not in TASK_STATUSES:
            raise ValueError("任务状态无效。")
        if status_value in {"completed", "closed"}:
            # 完成与关闭均要求无未闭环隐患，防止直接置 closed 绕过校验
            open_hazards = int(current and len([
                h for h in current.get("hazards", []) if h["status"] != "closed"
            ]))
            if open_hazards:
                raise ValueError("存在未闭环隐患，任务不能标记完成或关闭。")
        with self._lock, self._connect() as connection:
            connection.execute(
                "UPDATE patrol_tasks SET status = ?, updated_at = ? WHERE task_id = ?",
                (status_value, _now(), task_id),
            )
        return self.get_task(task_id, actor)

    def reopen_task(self, task_id: str, actor: dict[str, Any]) -> dict[str, Any]:
        """关闭/已完成/历史任务重启为执行中。"""
        current = self.get_task(task_id, actor)
        if not actor.get("is_admin"):
            raise ValueError("无权限重启任务。")
        if current["status"] not in {"completed", "closed"} and not current.get("legacy"):
            raise ValueError("只有已完成、已关闭或历史任务可以重启。")
        with self._lock, self._connect() as connection:
            connection.execute(
                "UPDATE patrol_tasks SET status = 'executing', legacy = 0, updated_at = ? WHERE task_id = ?",
                (_now(), task_id),
            )
        return self.get_task(task_id, actor)

    def soft_delete_task(self, task_id: str, actor: dict[str, Any]) -> dict[str, Any]:
        current = self.get_task(task_id, actor)
        if not actor.get("is_admin"):
            raise ValueError("无权限删除任务。")
        if current["status"] not in {"pending"}:
            raise ValueError("只有待执行的任务可以删除。")
        with self._lock, self._connect() as connection:
            connection.execute(
                "UPDATE patrol_tasks SET deleted = 1, updated_at = ? WHERE task_id = ?",
                (_now(), task_id),
            )
        return {"task_id": task_id, "task_no": current["task_no"], "deleted": True}

    def statistics(self, actor: dict[str, Any], *, line: str = "", date_from: str = "", date_to: str = "") -> dict[str, Any]:
        clauses = ["deleted = 0"]
        params: list[Any] = []
        if not actor.get("is_admin"):
            user_id = str(actor.get("user_id") or "")
            if not user_id:
                clauses.append("1 = 0")
            else:
                clauses.append(
                    "(EXISTS (SELECT 1 FROM patrol_task_users tu WHERE tu.task_id = patrol_tasks.task_id AND tu.user_id = ?) "
                    "OR assigned_user_id = ?)"
                )
                params.extend([user_id, user_id])
        if line.strip():
            clauses.append(
                "EXISTS (SELECT 1 FROM patrol_task_lines pl WHERE pl.task_id = patrol_tasks.task_id AND pl.line = ?)"
            )
            params.append(line.strip())
        if date_from.strip():
            clauses.append("substr(created_at, 1, 10) >= ?")
            params.append(date_from.strip())
        if date_to.strip():
            clauses.append("substr(created_at, 1, 10) <= ?")
            params.append(date_to.strip())
        where = "WHERE " + " AND ".join(clauses)
        counts = {key: 0 for key in TASK_STATUSES}
        counts["total"] = 0
        with self._connect() as connection:
            counts["total"] = int(connection.execute(
                f"SELECT COUNT(*) FROM patrol_tasks {where}", params
            ).fetchone()[0])
            rows = connection.execute(
                f"SELECT status, COUNT(*) AS c FROM patrol_tasks {where} GROUP BY status", params
            ).fetchall()
        for row in rows:
            counts[row["status"]] = int(row["c"])
        return counts

    # ---- 巡查记录与媒体 ----

    def create_record(self, task_id: str, data: dict[str, Any], actor: dict[str, Any]) -> dict[str, Any]:
        task = self.get_task(task_id, actor)
        if not self._visible(task, actor):
            raise KeyError("巡查任务不存在。")
        # 状态限制：已完成/已关闭的任务不允许再追加巡查记录
        if task["status"] not in {"pending", "executing"}:
            raise ValueError("任务已完成或关闭，不能再添加巡查记录。")
        record_type = _text(data.get("type") or "patrol", field="记录类型", required=True, maximum=20)
        if record_type not in RECORD_TYPES:
            raise ValueError("记录类型无效。")
        hazard_id = ""
        if record_type == "rectify":
            hazard_id = _text(data.get("hazard_id"), field="关联隐患", required=True, maximum=40)
            hazard = self.get_hazard(hazard_id)
            if hazard["task_id"] != task_id:
                raise ValueError("隐患不属于该任务。")
            if hazard["status"] not in {"pending_rectify", "rectifying"}:
                raise ValueError("该隐患当前状态不允许提交整改。")
        longitude = _optional_float(data.get("longitude"), field="经度", minimum=-180, maximum=180)
        latitude = _optional_float(data.get("latitude"), field="纬度", minimum=-90, maximum=90)
        accuracy = _optional_float(data.get("accuracy"), field="定位精度", minimum=0, maximum=100000)
        title = _text(data.get("title"), field="本次上传名称", maximum=120)
        note = _text(data.get("note"), field="备注", maximum=1000)
        record_id, now = _id("prec"), _now()
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO patrol_records(
                    record_id, task_id, type, hazard_id, title, longitude, latitude, accuracy,
                    note, created_by, created_by_name, created_at
                ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    record_id, task_id, record_type, hazard_id, title, longitude, latitude, accuracy,
                    note, actor.get("user_id") or "", actor.get("name") or "", now,
                ),
            )
            connection.execute(
                "UPDATE patrol_tasks SET updated_at = ? WHERE task_id = ?", (now, task_id)
            )
            # 上传整改反馈时，隐患进入整改中
            if record_type == "rectify":
                connection.execute(
                    "UPDATE patrol_hazards SET status = 'rectifying', updated_at = ? WHERE hazard_id = ?",
                    (now, hazard_id),
                )
        return self.get_record(record_id)

    def get_record(self, record_id: str) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM patrol_records WHERE record_id = ?", (record_id,)).fetchone()
            if row is None:
                raise KeyError("巡查记录不存在。")
            media = connection.execute(
                "SELECT * FROM patrol_media WHERE record_id = ? ORDER BY sort, created_at", (record_id,)
            ).fetchall()
        record = dict(row)
        record["media"] = [dict(item) for item in media]
        return record

    def update_record_fields(self, record_id: str, note: str | None = None, title: str | None = None, actor: dict[str, Any] | None = None) -> dict[str, Any]:
        record = self.get_record(record_id)
        if actor is None:
            actor = {}
        if not actor.get("is_admin") and not self._is_owner(actor, record.get("created_by") or ""):
            raise ValueError("无权限修改该记录。")
        # 状态限制：已完成/已关闭的任务不允许再修改记录
        task = self.get_task(record["task_id"], actor)
        if task["status"] not in {"pending", "executing"}:
            raise ValueError("任务已完成或关闭，不能再修改巡查记录。")
        note_value = _text(note if note is not None else record.get("note"), field="备注", maximum=1000)
        title_value = _text(title if title is not None else record.get("title"), field="本次上传名称", maximum=120)
        with self._lock, self._connect() as connection:
            connection.execute(
                "UPDATE patrol_records SET note = ?, title = ? WHERE record_id = ?",
                (note_value, title_value, record_id),
            )
        return self.get_record(record_id)

    def add_media(self, record_id: str, kind: str, stored_path: str | Path, file_name: str,
                  taken_at: str = "", caption: str = "") -> dict[str, Any]:
        if kind not in MEDIA_KINDS:
            raise ValueError("媒体类型无效。")
        record = self.get_record(record_id)
        now = _now()
        with self._lock, self._connect() as connection:
            counts = {
                row["kind"]: int(row["c"])
                for row in connection.execute(
                    "SELECT kind, COUNT(*) AS c FROM patrol_media WHERE record_id = ? GROUP BY kind",
                    (record_id,),
                )
            }
            photo_count = counts.get("photo", 0)
            video_count = counts.get("video", 0)
            if kind == "photo" and photo_count >= MAX_PHOTOS_PER_RECORD:
                raise ValueError(f"每条记录最多{MAX_PHOTOS_PER_RECORD}张照片。")
            if kind == "video" and video_count >= MAX_VIDEOS_PER_RECORD:
                raise ValueError(f"每条记录最多{MAX_VIDEOS_PER_RECORD}个视频。")
            sort = photo_count + video_count + 1
            media_id = _id("pmed")
            caption_value = _text(caption, field="照片名称", maximum=200)
            connection.execute(
                """
                INSERT INTO patrol_media(media_id, record_id, kind, file_name, file_path,
                                         caption, taken_at, sort, created_at)
                VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (media_id, record_id, kind, file_name, str(stored_path),
                 caption_value, taken_at, sort, now),
            )
        return self.get_media(media_id)

    def get_media(self, media_id: str) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM patrol_media WHERE media_id = ?", (media_id,)).fetchone()
        if row is None:
            raise KeyError("媒体文件不存在。")
        return dict(row)

    def media_file_path(self, media_id: str) -> Path:
        media = self.get_media(media_id)
        path = Path(media["file_path"]).resolve()
        if not path.exists() or not path.is_file():
            raise KeyError("媒体文件缺失。")
        return path

    def submit_rectify_review(self, hazard_id: str, actor: dict[str, Any]) -> dict[str, Any]:
        """巡查员提交整改复核（整改反馈上传完成后调用，隐患 → 待复核）。"""
        hazard = self.get_hazard(hazard_id)
        # 行级隔离：仅平台管理员或该隐患所属任务可见者可提交复核
        self.get_task(hazard["task_id"], actor)
        if hazard["status"] != "rectifying":
            raise ValueError("只有整改中的隐患可以提交复核。")
        # 必须有至少一条含整改图片的整改反馈记录（整改证据）
        with self._connect() as connection:
            evidence = int(connection.execute(
                """
                SELECT COUNT(*) AS c FROM patrol_media m
                JOIN patrol_records r ON r.record_id = m.record_id
                WHERE r.type = 'rectify' AND r.hazard_id = ? AND m.kind = 'photo'
                """,
                (hazard_id,),
            ).fetchone()["c"])
        if evidence <= 0:
            raise ValueError("请先上传整改图片，再提交整改完成。")
        now = _now()
        with self._lock, self._connect() as connection:
            # 条件写入防并发竞态：仅当仍处于整改中才允许转换
            cursor = connection.execute(
                """
                UPDATE patrol_hazards SET status = 'pending_review', submitter = ?, submit_time = ?, updated_at = ?
                WHERE hazard_id = ? AND status = 'rectifying'
                """,
                (actor.get("name") or "", now, now, hazard_id),
            )
            if cursor.rowcount == 0:
                raise ValueError("隐患状态已变更，请刷新后重试。")
        return self.get_hazard(hazard_id)

    # ---- 隐患 ----

    @staticmethod
    def _hazard(row: sqlite3.Row) -> dict[str, Any]:
        return dict(row)

    def get_hazard(self, hazard_id: str) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM patrol_hazards WHERE hazard_id = ?", (hazard_id,)).fetchone()
            if row is None:
                raise KeyError("问题隐患不存在。")
            shots = connection.execute(
                "SELECT * FROM patrol_hazard_shots WHERE hazard_id = ? ORDER BY created_at, shot_id",
                (hazard_id,),
            ).fetchall()
        hazard = self._hazard(row)
        hazard["shots"] = [dict(shot) for shot in shots]
        return hazard

    def add_hazard_shot(self, hazard_id: str, stored_path: str | Path, file_name: str) -> dict[str, Any]:
        self.get_hazard(hazard_id)
        shot_id, now = _id("pshot"), _now()
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO patrol_hazard_shots(shot_id, hazard_id, file_name, file_path, created_at)
                VALUES(?, ?, ?, ?, ?)
                """,
                (shot_id, hazard_id, file_name, str(stored_path), now),
            )
        return self.get_shot(shot_id)

    def get_shot(self, shot_id: str) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM patrol_hazard_shots WHERE shot_id = ?", (shot_id,)
            ).fetchone()
        if row is None:
            raise KeyError("隐患截图不存在。")
        return dict(row)

    def shot_file_path(self, shot_id: str) -> Path:
        shot = self.get_shot(shot_id)
        path = Path(shot["file_path"]).resolve()
        if not path.exists() or not path.is_file():
            raise KeyError("隐患截图文件缺失。")
        return path

    def delete_shot(self, shot_id: str, actor: dict[str, Any]) -> dict[str, Any]:
        shot = self.get_shot(shot_id)
        hazard = self.get_hazard(shot["hazard_id"])
        if not actor.get("is_admin") and not self._is_owner(actor, hazard.get("created_by") or ""):
            raise ValueError("无权限删除该截图。")
        # 状态限制：隐患进入复核或闭环后整改证据不可删除（管理员除外），保证审计链完整
        if not actor.get("is_admin") and hazard["status"] in {"pending_review", "closed"}:
            raise ValueError("隐患已提交复核或闭环，整改截图不可删除。")
        with self._lock, self._connect() as connection:
            connection.execute("DELETE FROM patrol_hazard_shots WHERE shot_id = ?", (shot_id,))
        Path(shot["file_path"]).unlink(missing_ok=True)
        return {"shot_id": shot_id, "deleted": True}

    # ---- 监测方案文档 ----

    def add_task_doc(
        self, task_id: str, file_name: str, stored_path: str | Path, kind: str, size: int, actor: dict[str, Any]
    ) -> dict[str, Any]:
        if not actor.get("is_admin"):
            raise ValueError("无权限上传监测方案文档。")
        task = self.get_task(task_id, actor)
        doc_id, now = _id("pdoc"), _now()
        with self._lock, self._connect() as connection:
            count = int(connection.execute(
                "SELECT COUNT(*) FROM patrol_task_docs WHERE task_id = ?", (task_id,)
            ).fetchone()[0])
            if count >= MAX_DOCS_PER_TASK:
                raise ValueError(f"每个任务最多{MAX_DOCS_PER_TASK}个监测方案文档。")
            connection.execute(
                """
                INSERT INTO patrol_task_docs(doc_id, task_id, file_name, file_path, kind, size, created_at)
                VALUES(?, ?, ?, ?, ?, ?, ?)
                """,
                (doc_id, task_id, file_name, str(stored_path), kind, size, now),
            )
        return self.get_task_doc(doc_id)

    def get_task_doc(self, doc_id: str) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM patrol_task_docs WHERE doc_id = ?", (doc_id,)).fetchone()
        if row is None:
            raise KeyError("监测方案文档不存在。")
        return dict(row)

    def doc_file_path(self, doc_id: str) -> Path:
        doc = self.get_task_doc(doc_id)
        path = Path(doc["file_path"]).resolve()
        if not path.exists() or not path.is_file():
            raise KeyError("监测方案文档文件缺失。")
        return path

    def delete_task_doc(self, doc_id: str, actor: dict[str, Any]) -> dict[str, Any]:
        if not actor.get("is_admin"):
            raise ValueError("无权限删除监测方案文档。")
        doc = self.get_task_doc(doc_id)
        with self._lock, self._connect() as connection:
            connection.execute("DELETE FROM patrol_task_docs WHERE doc_id = ?", (doc_id,))
        Path(doc["file_path"]).unlink(missing_ok=True)
        return {"doc_id": doc_id, "deleted": True}

    @staticmethod
    def _is_owner(actor: dict[str, Any], created_by: str) -> bool:
        """行级归属判定：空身份或空归属不得因两侧同为空串而判通过。"""
        user_id = str(actor.get("user_id") or "")
        return bool(user_id) and created_by == user_id

    def update_hazard(self, hazard_id: str, data: dict[str, Any], actor: dict[str, Any]) -> dict[str, Any]:
        current = self.get_hazard(hazard_id)
        if not actor.get("is_admin") and not self._is_owner(actor, current.get("created_by") or ""):
            raise ValueError("无权限修改该隐患。")
        # 状态限制：隐患被平台确认进入整改流程后仅管理员可修改，保证闭环数据完整
        if not actor.get("is_admin") and current["status"] != "pending_confirm":
            raise ValueError("隐患已被确认，仅平台管理员可以修改。")
        description = _text(data.get("description", current["description"]), field="隐患描述", required=True, maximum=2000)
        hazard_type = _text(data.get("hazard_type", current["hazard_type"]), field="隐患类型", maximum=60)
        risk_level = _text(data.get("risk_level", current["risk_level"]), field="风险等级", maximum=40)
        rectify_owner = _text(data.get("rectify_owner", current["rectify_owner"]), field="整改责任人", maximum=120)
        rectify_requirement = _text(data.get("rectify_requirement", current["rectify_requirement"]), field="整改要求", maximum=2000)
        if not actor.get("is_admin"):
            # 平台字段保护：巡查员不得篡改平台下发的整改要求与责任人
            rectify_owner = current["rectify_owner"]
            rectify_requirement = current["rectify_requirement"]
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                UPDATE patrol_hazards SET description = ?, hazard_type = ?, risk_level = ?,
                    rectify_owner = ?, rectify_requirement = ?, updated_at = ?
                WHERE hazard_id = ?
                """,
                (description, hazard_type, risk_level, rectify_owner, rectify_requirement, _now(), hazard_id),
            )
        return self.get_hazard(hazard_id)

    def delete_hazard(self, hazard_id: str, actor: dict[str, Any]) -> dict[str, Any]:
        current = self.get_hazard(hazard_id)
        if not actor.get("is_admin") and not self._is_owner(actor, current.get("created_by") or ""):
            raise ValueError("无权限删除该隐患。")
        # 状态限制：隐患一旦被确认并进入整改流程后不可删除，保证闭环数据完整
        if current["status"] != "pending_confirm":
            raise ValueError("只有待确认状态的隐患可以删除。")
        with self._lock, self._connect() as connection:
            shots = connection.execute(
                "SELECT file_path FROM patrol_hazard_shots WHERE hazard_id = ?", (hazard_id,)
            ).fetchall()
            # 条件删除防并发竞态：仅当仍处于待确认状态才允许删除
            cursor = connection.execute(
                "DELETE FROM patrol_hazards WHERE hazard_id = ? AND status = 'pending_confirm'", (hazard_id,)
            )
            if cursor.rowcount == 0:
                raise ValueError("隐患状态已变更，请刷新后重试。")
            connection.execute("DELETE FROM patrol_hazard_shots WHERE hazard_id = ?", (hazard_id,))
        for shot in shots:
            Path(shot["file_path"]).unlink(missing_ok=True)
        return {"hazard_id": hazard_id, "deleted": True}

    def create_hazard(self, task_id: str, data: dict[str, Any], actor: dict[str, Any]) -> dict[str, Any]:
        task = self.get_task(task_id, actor)
        if not self._visible(task, actor):
            raise KeyError("巡查任务不存在。")
        # 状态限制：已完成/已关闭的任务不允许再新增隐患，避免破坏"完成即无未闭环隐患"不变量
        if task["status"] not in {"pending", "executing"}:
            raise ValueError("任务已完成或关闭，不能再新增隐患。")
        description = _text(data.get("description"), field="隐患描述", required=True, maximum=2000)
        hazard_type = _text(data.get("hazard_type"), field="隐患类型", maximum=60)
        risk_level = _text(data.get("risk_level"), field="风险等级", maximum=40)
        record_id = _text(data.get("record_id"), field="关联巡查记录", maximum=40)
        # record_id 归属校验：防止跨任务悬挂引用
        if record_id:
            record = self.get_record(record_id)
            if record["task_id"] != task_id:
                raise ValueError("关联巡查记录不属于该任务。")
        rectify_owner = _text(data.get("rectify_owner"), field="整改责任人", maximum=120)
        is_admin = actor.get("is_admin")
        # 平台记录 → 直接待整改 + 整改要求；巡查员记录 → 待确认
        if is_admin:
            status_value = "pending_rectify"
            rectify_requirement = _text(data.get("rectify_requirement"), field="整改要求", required=True, maximum=2000)
            created_by_role = "platform"
        else:
            status_value = "pending_confirm"
            rectify_requirement = ""
            created_by_role = "patrol"
        hazard_id, now = _id("phaz"), _now()
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                INSERT INTO patrol_hazards(
                    hazard_id, task_id, record_id, description, hazard_type, risk_level,
                    rectify_requirement, rectify_owner, status, created_by, created_by_name,
                    created_by_role, created_at, updated_at
                ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    hazard_id, task_id, record_id, description, hazard_type, risk_level,
                    rectify_requirement, rectify_owner, status_value, actor.get("user_id") or "",
                    actor.get("name") or "", created_by_role, now, now,
                ),
            )
            connection.execute(
                "UPDATE patrol_tasks SET updated_at = ? WHERE task_id = ?", (now, task_id)
            )
        return self.get_hazard(hazard_id)

    def confirm_hazard(self, hazard_id: str, data: dict[str, Any], actor: dict[str, Any]) -> dict[str, Any]:
        """平台确认隐患并下发整改要求（待确认 → 待整改）。"""
        if not actor.get("is_admin"):
            raise ValueError("无权限确认隐患。")
        current = self.get_hazard(hazard_id)
        if current["status"] != "pending_confirm":
            raise ValueError("只有待确认的隐患可以确认。")
        rectify_requirement = _text(data.get("rectify_requirement"), field="整改要求", required=True, maximum=2000)
        now = _now()
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                UPDATE patrol_hazards SET rectify_requirement = ?, status = 'pending_rectify',
                    confirmer = ?, confirm_time = ?, updated_at = ?
                WHERE hazard_id = ?
                """,
                (rectify_requirement, actor.get("name") or "", now, now, hazard_id),
            )
        return self.get_hazard(hazard_id)

    def review_hazard(self, hazard_id: str, data: dict[str, Any], actor: dict[str, Any]) -> dict[str, Any]:
        """平台复核整改（待复核 → 已闭环 / 退回整改中）。"""
        if not actor.get("is_admin"):
            raise ValueError("无权限复核隐患。")
        current = self.get_hazard(hazard_id)
        if current["status"] != "pending_review":
            raise ValueError("只有待复核的隐患可以复核。")
        result = _text(data.get("result"), field="复核结论", required=True, maximum=20)
        if result not in {"closed", "reject"}:
            raise ValueError("复核结论无效。")
        comment = _text(data.get("comment"), field="复核意见", maximum=2000)
        now = _now()
        new_status = "closed" if result == "closed" else "rectifying"
        with self._lock, self._connect() as connection:
            connection.execute(
                """
                UPDATE patrol_hazards SET status = ?, reviewer = ?, review_comment = ?,
                    review_time = ?, updated_at = ?
                WHERE hazard_id = ?
                """,
                (new_status, actor.get("name") or "", comment, now, now, hazard_id),
            )
        return self.get_hazard(hazard_id)


# ---------------------------------------------------------------------------
# 水印
# ---------------------------------------------------------------------------

_CJK_FONT_CANDIDATES = (
    "C:/Windows/Fonts/msyh.ttc",
    "C:/Windows/Fonts/msyh.ttf",
    "C:/Windows/Fonts/simhei.ttf",
    "C:/Windows/Fonts/simsun.ttc",
    "C:/Windows/Fonts/Deng.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
)


def _load_cjk_font(size: int):
    try:
        from PIL import ImageFont
    except ImportError:
        return None
    for candidate in _CJK_FONT_CANDIDATES:
        if Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size)
            except Exception:
                continue
    try:
        return ImageFont.load_default()
    except Exception:
        return None


def apply_watermark(source: Path, destination: Path, lines: list[str]) -> Path:
    """在图片底部叠加半透明水印，返回目标路径。失败时退回原图复制。"""
    try:
        from PIL import Image, ImageDraw, ImageOps
    except ImportError:
        _copy_file(source, destination)
        return destination
    try:
        with Image.open(source) as opened:
            image = ImageOps.exif_transpose(opened)
            image = image.convert("RGB")
            width, height = image.size
            font_size = max(18, min(42, width // 26))
            font = _load_cjk_font(font_size)
            draw = ImageDraw.Draw(image, "RGBA")
            line_height = int(font_size * 1.4)
            band_height = line_height * len(lines) + 20
            if band_height < height:
                draw.rectangle([(0, height - band_height), (width, height)], fill=(0, 0, 0, 130))
            y = height - band_height + 10
            for line in lines:
                draw.text((16, y), line, font=font, fill=(255, 255, 255, 240))
                y += line_height
            image.save(destination, "JPEG", quality=90)
        return destination
    except Exception as exc:  # noqa: BLE001
        LOGGER.warning("水印失败，退回原图：%s", exc)
        _copy_file(source, destination)
        return destination


def _copy_file(source: Path, destination: Path) -> None:
    import shutil

    shutil.copyfile(source, destination)


# ---------------------------------------------------------------------------
# 认证与路由
# ---------------------------------------------------------------------------


def verify_patrol_access(
    x_service_token: Annotated[str | None, Header()] = None,
    x_dev_token: Annotated[str | None, Header()] = None,
) -> None:
    service_ok = (not SERVICE_TOKEN) or hmac.compare_digest(x_service_token or "", SERVICE_TOKEN)
    if service_ok:
        return
    if PATROL_DEV_TOKEN and hmac.compare_digest(x_dev_token or "", PATROL_DEV_TOKEN):
        return
    raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="无访问权限。")


def get_actor(
    x_actor_user_id: Annotated[str, Header()] = "",
    x_actor_name: Annotated[str, Header()] = "",
    x_actor_is_admin: Annotated[str, Header()] = "",
) -> dict[str, Any]:
    user_id = x_actor_user_id or ""
    return {
        "user_id": user_id,
        "name": x_actor_name or user_id,
        # 仅当 Java 代理明确注入 X-Actor-Is-Admin=true 时才视为平台管理员；
        # 缺失身份头（user_id 为空）不再默认提权，避免越权风险。
        "is_admin": x_actor_is_admin.lower() in {"1", "true", "yes"},
    }


_repository: PatrolRepository | None = None


def get_repository() -> PatrolRepository:
    global _repository
    if _repository is None:
        _repository = PatrolRepository()
    return _repository


router = APIRouter(prefix="/api/v1/patrol", tags=["patrol"], dependencies=[Depends(verify_patrol_access)])


class PatrolDictCreatePayload(BaseModel):
    type: str = Field(min_length=1, max_length=40)
    label: str = Field(min_length=1, max_length=60)
    value: str = Field(default="", max_length=80)
    sort: int = Field(default=0, ge=0, le=10000)


class PatrolDictUpdatePayload(BaseModel):
    label: str | None = Field(default=None, min_length=1, max_length=60)
    value: str | None = Field(default=None, max_length=80)
    sort: int | None = Field(default=None, ge=0, le=10000)
    enabled: int | None = Field(default=None, ge=0, le=1)


class PatrolTaskCreatePayload(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    lines: list[str] = Field(default_factory=list)
    user_ids: list[str] = Field(default_factory=list)
    user_names: list[str] = Field(default_factory=list)
    line: str = Field(default="", max_length=200)
    location_desc: str = Field(default="", max_length=200)
    requirement: str = Field(default="", max_length=2000)
    assigned_user_id: str = Field(default="", max_length=40)
    assigned_user_name: str = Field(default="", max_length=60)
    remark: str = Field(default="", max_length=1000)
    monitor_frequency: str = Field(default="", max_length=2000)
    monitor_points: str = Field(default="", max_length=2000)
    warning_threshold: str = Field(default="", max_length=2000)
    emergency_plan: str = Field(default="", max_length=2000)
    report_requirement: str = Field(default="", max_length=2000)
    review_opinion: str = Field(default="", max_length=2000)


class PatrolWorkflowTaskPayload(BaseModel):
    workflow_id: str = Field(min_length=1, max_length=80)
    project_id: str = Field(default="", max_length=80)
    stage_id: str = Field(default="", max_length=80)
    stage_name: str = Field(default="", max_length=100)
    project_name: str = Field(min_length=1, max_length=120)
    initiator_id: str | int | None = None
    initiator_name: str = Field(default="", max_length=60)
    latest_summary: str = Field(default="", max_length=2000)
    latest_result_json: str | dict[str, Any] | None = None
    final_approval_opinion: str = Field(default="", max_length=500)
    # D122：终审通过时提交的创建参数（可选；未传则回退指派发起人，传空数组=不指派）
    user_ids: list[str] | None = None
    user_names: list[str] | None = None
    location_desc: str = Field(default="", max_length=200)
    remark: str = Field(default="", max_length=1000)


class PatrolTaskUpdatePayload(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=120)
    lines: list[str] | None = None
    user_ids: list[str] | None = None
    user_names: list[str] | None = None
    line: str | None = Field(default=None, max_length=200)
    location_desc: str | None = Field(default=None, max_length=200)
    requirement: str | None = Field(default=None, max_length=2000)
    assigned_user_id: str | None = Field(default=None, max_length=200)
    assigned_user_name: str | None = Field(default=None, max_length=200)
    monitor_frequency: str | None = Field(default=None, max_length=2000)
    monitor_points: str | None = Field(default=None, max_length=2000)
    warning_threshold: str | None = Field(default=None, max_length=2000)
    emergency_plan: str | None = Field(default=None, max_length=2000)
    report_requirement: str | None = Field(default=None, max_length=2000)
    review_opinion: str | None = Field(default=None, max_length=2000)


class PatrolTaskStatusPayload(BaseModel):
    status: Literal["pending", "executing", "completed", "closed"]


class PatrolRecordCreatePayload(BaseModel):
    type: Literal["patrol", "rectify"]
    hazard_id: str = Field(default="", max_length=40)
    title: str = Field(default="", max_length=120)
    longitude: float | None = Field(default=None, ge=-180, le=180)
    latitude: float | None = Field(default=None, ge=-90, le=90)
    accuracy: float | None = Field(default=None, ge=0, le=100000)
    note: str = Field(default="", max_length=1000)


class PatrolRecordUpdatePayload(BaseModel):
    note: str | None = Field(default=None, max_length=1000)
    title: str | None = Field(default=None, max_length=120)


class PatrolHazardCreatePayload(BaseModel):
    description: str = Field(min_length=1, max_length=2000)
    hazard_type: str = Field(default="", max_length=60)
    risk_level: str = Field(default="", max_length=40)
    record_id: str = Field(default="", max_length=40)
    rectify_owner: str = Field(default="", max_length=120)
    rectify_requirement: str = Field(default="", max_length=2000)


class PatrolHazardUpdatePayload(BaseModel):
    description: str = Field(min_length=1, max_length=2000)
    hazard_type: str = Field(default="", max_length=60)
    risk_level: str = Field(default="", max_length=40)
    rectify_owner: str = Field(default="", max_length=120)
    rectify_requirement: str = Field(default="", max_length=2000)


class PatrolHazardConfirmPayload(BaseModel):
    rectify_requirement: str = Field(min_length=1, max_length=2000)


class PatrolHazardReviewPayload(BaseModel):
    result: Literal["closed", "reject"]
    comment: str = Field(default="", max_length=2000)


class PatrolOpinionReviewPayload(BaseModel):
    result: Literal["pass", "return"]
    comment: str = Field(default="", max_length=2000)


class PatrolOpinionCheckPayload(BaseModel):
    comment: str = Field(max_length=2000)


def _model_values(payload: BaseModel) -> dict[str, Any]:
    if hasattr(payload, "model_dump"):
        return payload.model_dump(exclude_unset=True)
    return payload.dict(exclude_unset=True)


@router.get("/dicts")
def list_dicts(dict_type: str = "line", repo: PatrolRepository = Depends(get_repository)) -> list[dict[str, Any]]:
    try:
        return repo.list_dicts(dict_type)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


def _require_admin(actor: dict) -> None:
    """字典等全局配置的写操作仅允许平台管理员。"""
    if not actor.get("is_admin"):
        raise HTTPException(status_code=403, detail="仅平台管理员可维护字典。")


@router.post("/dicts", status_code=201)
def create_dict(payload: PatrolDictCreatePayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    _require_admin(actor)
    try:
        return repo.create_dict(_model_values(payload))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/dicts/{dict_id}")
def update_dict(dict_id: str, payload: PatrolDictUpdatePayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    _require_admin(actor)
    try:
        return repo.update_dict(dict_id, _model_values(payload))
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.delete("/dicts/{dict_id}")
def delete_dict(dict_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    _require_admin(actor)
    try:
        return repo.delete_dict(dict_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@router.post("/tasks", status_code=201)
def create_task(payload: PatrolTaskCreatePayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.create_task(_model_values(payload), actor)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/workflow-tasks", status_code=201)
def create_workflow_task(
    payload: PatrolWorkflowTaskPayload,
    repo: PatrolRepository = Depends(get_repository),
) -> dict[str, Any]:
    try:
        return repo.create_workflow_task(_model_values(payload))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.get("/tasks")
def list_tasks(
    page: int = 1,
    size: int = 20,
    line: str = "",
    status: str = "",
    task_type: str = "",
    assigned_user_id: str = "",
    date_from: str = "",
    date_to: str = "",
    keyword: str = "",
    sort: str = "created",
    actor: dict = Depends(get_actor),
    repo: PatrolRepository = Depends(get_repository),
) -> dict[str, Any]:
    return repo.list_tasks(
        actor,
        page=page, size=size, line=line, status_value=status, task_type_value=task_type,
        assigned_user_id=assigned_user_id, date_from=date_from, date_to=date_to,
        keyword=keyword, order_by=sort if sort in {"created", "recent"} else "created",
    )


@router.get("/statistics")
def statistics(
    line: str = "",
    date_from: str = "",
    date_to: str = "",
    actor: dict = Depends(get_actor),
    repo: PatrolRepository = Depends(get_repository),
) -> dict[str, Any]:
    return repo.statistics(actor, line=line, date_from=date_from, date_to=date_to)


@router.get("/tasks/{task_id}")
def get_task(task_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.get_task(task_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@router.get("/projects/{project_id}/task")
def get_project_task(project_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any] | None:
    """按项目档案查唯一关联巡查任务（供档案页跳转；不存在/软删返回 null）。"""
    return repo.get_project_task(project_id, actor)


@router.post("/tasks/{task_id}")
def update_task(task_id: str, payload: PatrolTaskUpdatePayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.update_task(task_id, _model_values(payload), actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/tasks/{task_id}/status")
def set_task_status(task_id: str, payload: PatrolTaskStatusPayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.set_task_status(task_id, payload.status, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/tasks/{task_id}/reopen")
def reopen_task(task_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.reopen_task(task_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.delete("/tasks/{task_id}")
def soft_delete_task(task_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.soft_delete_task(task_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@router.post("/tasks/{task_id}/records", status_code=201)
def create_record(task_id: str, payload: PatrolRecordCreatePayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.create_record(task_id, _model_values(payload), actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/records/{record_id}")
def update_record(record_id: str, payload: PatrolRecordUpdatePayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.update_record_fields(record_id, note=payload.note, title=payload.title, actor=actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


async def _save_media(upload: UploadFile, destination: Path, maximum: int, allowed: set[str]) -> None:
    suffix = Path(upload.filename or "").suffix.lower()
    if suffix not in allowed:
        raise HTTPException(status_code=415, detail=f"不支持的文件类型：{suffix or '无扩展名'}")
    total = 0
    try:
        with destination.open("wb") as target:
            while chunk := await upload.read(1024 * 1024):
                total += len(chunk)
                if total > maximum:
                    raise HTTPException(status_code=413, detail="文件大小超过限制。")
                target.write(chunk)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    finally:
        await upload.close()


@router.post("/records/{record_id}/media", status_code=201)
async def add_media(
    record_id: str,
    file: UploadFile = File(...),
    kind: str = Form("photo"),
    taken_at: str = Form(""),
    caption: str = Form(""),
    actor: dict = Depends(get_actor),
    repo: PatrolRepository = Depends(get_repository),
) -> dict[str, Any]:
    if kind not in MEDIA_KINDS:
        raise HTTPException(status_code=422, detail="媒体类型无效。")
    try:
        record = repo.get_record(record_id)
        # 行级隔离：校验记录所属任务对当前用户可见（巡查员只能给自己被指派的任务上传媒体）
        task = repo.get_task(record["task_id"], actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    # 状态限制：已完成/已关闭的任务不允许再上传媒体
    if task["status"] not in {"pending", "executing"}:
        raise HTTPException(status_code=422, detail="任务已完成或关闭，不能再上传媒体。")
    suffix = Path(file.filename or (".jpg" if kind == "photo" else ".mp4")).suffix.lower()
    allowed = PHOTO_SUFFIXES if kind == "photo" else VIDEO_SUFFIXES
    if suffix not in allowed:
        suffix = ".jpg" if kind == "photo" else ".mp4"
    maximum = MAX_PHOTO_BYTES if kind == "photo" else MAX_VIDEO_BYTES
    folder = repo.upload_root / record_id
    folder.mkdir(parents=True, exist_ok=True)
    source = folder / f"{uuid.uuid4().hex}{suffix}"
    await _save_media(file, source, maximum, allowed)

    media_id = _id("pmed")
    destination = folder / f"{media_id}{suffix}"
    if kind == "photo":
        lines = [
            f"任务 {task['task_no']}",
            f"时间 {taken_at or _now()}",
            f"上报 {actor.get('name') or actor.get('user_id') or ''}".rstrip(),
        ]
        if task.get("line"):
            lines.append(f"线路 {task['line']}")
        if task.get("location_desc"):
            lines.append(f"位置 {task['location_desc'][:40]}")
        apply_watermark(source, destination, lines)
        source.unlink(missing_ok=True)
    else:
        destination = source

    try:
        return repo.add_media(record_id, kind, destination, Path(file.filename or destination.name).name, taken_at, caption)
    except ValueError as exc:
        destination.unlink(missing_ok=True)
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/tasks/{task_id}/hazards", status_code=201)
def create_hazard(task_id: str, payload: PatrolHazardCreatePayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.create_hazard(task_id, _model_values(payload), actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/hazards/{hazard_id}")
def update_hazard(hazard_id: str, payload: PatrolHazardUpdatePayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.update_hazard(hazard_id, _model_values(payload), actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.delete("/hazards/{hazard_id}")
def delete_hazard(hazard_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.delete_hazard(hazard_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.delete("/shots/{shot_id}")
def delete_shot(shot_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.delete_shot(shot_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/hazards/{hazard_id}/confirm")
def confirm_hazard(hazard_id: str, payload: PatrolHazardConfirmPayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.confirm_hazard(hazard_id, _model_values(payload), actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/hazards/{hazard_id}/submit")
def submit_rectify_review(hazard_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.submit_rectify_review(hazard_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/hazards/{hazard_id}/review")
def review_hazard(hazard_id: str, payload: PatrolHazardReviewPayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.review_hazard(hazard_id, _model_values(payload), actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.get("/media/{media_id}/file")
def media_file(media_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> FileResponse:
    try:
        media = repo.get_media(media_id)
        record = repo.get_record(media["record_id"])
        # 行级隔离：校验媒体所属任务对当前用户可见
        repo.get_task(record["task_id"], actor)
        path = repo.media_file_path(media_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    media_type = {
        ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
        ".webp": "image/webp", ".bmp": "image/bmp",
        ".mp4": "video/mp4", ".mov": "video/quicktime",
    }.get(path.suffix.lower(), "application/octet-stream")
    return FileResponse(path, media_type=media_type, filename=path.name)


@router.post("/hazards/{hazard_id}/shots", status_code=201)
async def add_hazard_shot(
    hazard_id: str,
    file: UploadFile = File(...),
    actor: dict = Depends(get_actor),
    repo: PatrolRepository = Depends(get_repository),
) -> dict[str, Any]:
    try:
        hazard = repo.get_hazard(hazard_id)
        # 行级隔离：仅平台管理员或隐患创建者可为该隐患上传整改截图
        if not actor.get("is_admin") and hazard.get("created_by") != str(actor.get("user_id") or ""):
            raise HTTPException(status_code=403, detail="无权限上传该隐患的整改截图。")
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    suffix = Path(file.filename or ".jpg").suffix.lower()
    if suffix not in PHOTO_SUFFIXES:
        suffix = ".jpg"
    folder = repo.upload_root / "shots" / hazard_id
    folder.mkdir(parents=True, exist_ok=True)
    source = folder / f"{uuid.uuid4().hex}{suffix}"
    await _save_media(file, source, MAX_PHOTO_BYTES, PHOTO_SUFFIXES)
    shot_id = _id("pshot")
    destination = folder / f"{shot_id}{suffix}"
    apply_watermark(source, destination, [
        f"隐患截图 {_now()}",
        f"上报 {actor.get('name') or actor.get('user_id') or ''}".rstrip(),
    ])
    source.unlink(missing_ok=True)
    try:
        return repo.add_hazard_shot(hazard_id, destination, Path(file.filename or destination.name).name)
    except KeyError as exc:
        destination.unlink(missing_ok=True)
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@router.get("/shots/{shot_id}/file")
def shot_file(shot_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> FileResponse:
    try:
        shot = repo.get_shot(shot_id)
        hazard = repo.get_hazard(shot["hazard_id"])
        # 行级隔离：校验截图所属隐患的任务对当前用户可见
        repo.get_task(hazard["task_id"], actor)
        path = repo.shot_file_path(shot_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    media_type = {
        ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
        ".webp": "image/webp", ".bmp": "image/bmp",
    }.get(path.suffix.lower(), "application/octet-stream")
    return FileResponse(path, media_type=media_type, filename=path.name)


# ---- 监测方案文档 ----

_DOC_MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
    ".webp": "image/webp", ".bmp": "image/bmp",
}


@router.post("/tasks/{task_id}/docs", status_code=201)
async def add_task_doc(
    task_id: str,
    file: UploadFile = File(...),
    actor: dict = Depends(get_actor),
    repo: PatrolRepository = Depends(get_repository),
) -> dict[str, Any]:
    suffix = Path(file.filename or ".pdf").suffix.lower()
    if suffix not in DOC_SUFFIXES:
        raise HTTPException(status_code=415, detail="仅支持 PDF、Word、图片文件。")
    if suffix in DOC_PDF_SUFFIXES:
        kind = "pdf"
    elif suffix in DOC_WORD_SUFFIXES:
        kind = "word"
    else:
        kind = "image"
    try:
        repo.get_task(task_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    folder = repo.upload_root / "docs" / task_id
    folder.mkdir(parents=True, exist_ok=True)
    destination = folder / f"{uuid.uuid4().hex}{suffix}"
    await _save_media(file, destination, MAX_DOC_BYTES, DOC_SUFFIXES)
    try:
        return repo.add_task_doc(task_id, Path(file.filename or destination.name).name, destination, kind, destination.stat().st_size, actor)
    except ValueError as exc:
        destination.unlink(missing_ok=True)
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.get("/docs/{doc_id}/file")
def doc_file(doc_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> FileResponse:
    try:
        doc = repo.get_task_doc(doc_id)
        # 行级隔离：校验文档所属任务对当前用户可见
        repo.get_task(doc["task_id"], actor)
        path = repo.doc_file_path(doc_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    media_type = _DOC_MEDIA_TYPES.get(path.suffix.lower(), "application/octet-stream")
    return FileResponse(path, media_type=media_type, filename=path.name)


@router.delete("/docs/{doc_id}")
def delete_task_doc(doc_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.delete_task_doc(doc_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


# ---- 审核意见 → 现场核查（Step①）----

@router.get("/tasks/{task_id}/opinions")
def list_task_opinions(task_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> list[dict[str, Any]]:
    try:
        return repo.list_task_opinions(task_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/tasks/{task_id}/opinions/sync")
def sync_task_opinions(task_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    """把项目档案中该项目意见同步进任务（幂等，供平台手动触发）。"""
    try:
        task = repo.get_task(task_id, actor)
        project_id = task.get("source_project_id") or ""
        if not project_id:
            raise ValueError("该任务未关联项目档案，无需同步。")
        added = repo.sync_project_opinions(task_id, project_id)
        return {"added": added, "total": len(repo.list_task_opinions(task_id, actor))}
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/opinions/{opinion_id}/photos", status_code=201)
async def add_opinion_photo(
    opinion_id: str,
    file: UploadFile = File(...),
    taken_at: str = Form(""),
    actor: dict = Depends(get_actor),
    repo: PatrolRepository = Depends(get_repository),
) -> dict[str, Any]:
    try:
        opinion = repo.get_opinion(opinion_id)
        repo.get_task(opinion["task_id"], actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    suffix = Path(file.filename or ".jpg").suffix.lower()
    if suffix not in PHOTO_SUFFIXES:
        raise HTTPException(status_code=415, detail="仅支持图片文件。")
    folder = repo.upload_root / "opinions" / opinion_id
    folder.mkdir(parents=True, exist_ok=True)
    source = folder / f"{uuid.uuid4().hex}{suffix}"
    await _save_media(file, source, MAX_OPINION_PHOTO_BYTES, PHOTO_SUFFIXES)
    photo_id = _id("pphoto")
    destination = folder / f"{photo_id}{suffix}"
    try:
        lines = [f"现场核查 {_now()}", (actor.get("name") or actor.get("user_id") or "")]
        apply_watermark(source, destination, lines)
        source.unlink(missing_ok=True)
        return repo.add_opinion_photo(
            opinion_id, destination, Path(file.filename or destination.name).name, taken_at, actor
        )
    except ValueError as exc:
        destination.unlink(missing_ok=True)
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.get("/opinion-photos/{photo_id}/file")
def opinion_photo_file(photo_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> FileResponse:
    try:
        with repo._connect() as connection:
            row = connection.execute(
                "SELECT * FROM patrol_opinion_photos WHERE photo_id = ?", (photo_id,)
            ).fetchone()
        if row is None:
            raise KeyError("意见照片不存在。")
        opinion = repo.get_opinion(row["opinion_id"])
        repo.get_task(opinion["task_id"], actor)  # 行级隔离
        path = repo.opinion_photo_path(photo_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    media_type = {
        ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
        ".webp": "image/webp", ".bmp": "image/bmp",
    }.get(path.suffix.lower(), "application/octet-stream")
    return FileResponse(path, media_type=media_type, filename=path.name)


@router.delete("/opinion-photos/{photo_id}")
def delete_opinion_photo(photo_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    try:
        return repo.delete_opinion_photo(photo_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/opinions/{opinion_id}/submit")
def submit_opinion(opinion_id: str, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    """巡查员：已有现场照片后提交，意见进入「已拍照（待核查）」。"""
    try:
        return repo.submit_opinion(opinion_id, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/opinions/{opinion_id}/review")
def review_opinion(opinion_id: str, payload: PatrolOpinionReviewPayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    """平台：通过=完成 / 退回重拍（照片保留）。"""
    try:
        return repo.review_opinion(opinion_id, _model_values(payload), actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/opinions/{opinion_id}/check")
def update_opinion_check(opinion_id: str, payload: PatrolOpinionCheckPayload, actor: dict = Depends(get_actor), repo: PatrolRepository = Depends(get_repository)) -> dict[str, Any]:
    """平台：编辑该意见项的核查意见（不改状态，供导出前维护）。"""
    try:
        return repo.update_opinion_check(opinion_id, payload.comment, actor)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


# ---------------------------------------------------------------------------
# 终审意见与现场核查 Word 导出（python-docx）
# ---------------------------------------------------------------------------

class WorkflowWordExportPayload(BaseModel):
    workflow_id: str = Field(default="", max_length=40)
    project_id: str = Field(default="", max_length=80)
    project_name: str = Field(default="", max_length=200)
    stage_name: str = Field(default="", max_length=100)
    audit_version: int = Field(default=1, ge=1)
    workflow_status: str = Field(default="", max_length=32)
    initiator_name: str = Field(default="", max_length=64)
    latest_summary: str = Field(default="", max_length=4000)
    latest_risk_level: str = Field(default="", max_length=16)
    final_opinion: str = Field(default="", max_length=2000)
    approved_time: str = Field(default="", max_length=40)
    logs: list[dict[str, Any]] = Field(default_factory=list)


_OP_STATUS_LABEL = {"pending_photo": "待拍照", "photo_taken": "已拍照·待核查", "done": "已完成", "returned": "退回重拍"}


def _patrol_check_block(project_id: str) -> dict[str, Any]:
    """按项目读取施工前核查任务及其逐条意见+现场照片，用于并入 Word。"""
    result: dict[str, Any] = {"task_no": "", "opinions": []}
    if not project_id:
        return result
    db = Path(PATROL_DB)
    if not db.exists():
        return result
    connection = sqlite3.connect(db, timeout=10)
    try:
        connection.row_factory = sqlite3.Row
        task = connection.execute(
            "SELECT task_id, task_no FROM patrol_tasks "
            "WHERE source_project_id = ? AND deleted = 0 ORDER BY created_at LIMIT 1",
            (project_id,),
        ).fetchone()
        if task is None:
            return result
        result["task_no"] = task["task_no"]
        opinions = connection.execute(
            "SELECT * FROM patrol_task_opinions WHERE task_id = ? ORDER BY sort, created_at, opinion_id",
            (task["task_id"],),
        ).fetchall()
        for row in opinions:
            photos = connection.execute(
                "SELECT file_path, file_name, created_at FROM patrol_opinion_photos "
                "WHERE opinion_id = ? ORDER BY sort, created_at", (row["opinion_id"],)
            ).fetchall()
            photo_paths = [Path(p["file_path"]) for p in photos
                           if Path(p["file_path"]).is_file()]
            result["opinions"].append({
                "title": row["title"] or "",
                "risk_level": row["risk_level"] or "",
                "stage_name": row["stage_name"] or "",
                "content": row["opinion_content"] or "",
                "status": row["status"],
                "check_opinion": row["check_opinion"] or "",
                "check_time": row["check_time"] or "",
                "return_reason": row["return_reason"] or "",
                "photos": photo_paths,
            })
    finally:
        connection.close()
    return result


def _build_workflow_word(payload: WorkflowWordExportPayload, block: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Inches, Pt

    doc = Document()
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal'].font.size = Pt(10.5)

    doc.add_heading('终审意见与现场核查记录', level=0)

    table = doc.add_table(rows=0, cols=4)
    table.style = 'Table Grid'
    info = [
        ('项目名称', payload.project_name, '阶段', payload.stage_name),
        ('流程编号', payload.workflow_id, '审核版本', str(payload.audit_version)),
        ('流程状态', payload.workflow_status, '发起人', payload.initiator_name),
        ('终审通过时间', payload.approved_time or '-', '对应巡查任务', block['task_no'] or '-'),
    ]
    for r in info:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v or '-')
    doc.add_paragraph()

    h = doc.add_heading('一、终审意见', level=1)
    risk = payload.latest_risk_level or '-'
    p = doc.add_paragraph()
    p.add_run('风险等级：').bold = True
    p.add_run(str(risk))
    p = doc.add_paragraph()
    p.add_run('终审综合结论：').bold = True
    p.add_run(str(payload.latest_summary or '无'))
    if payload.final_opinion:
        p = doc.add_paragraph()
        p.add_run('终审意见：').bold = True
        p.add_run(str(payload.final_opinion))

    if payload.logs:
        doc.add_heading('流转留痕', level=2)
        lt = doc.add_table(rows=1, cols=5)
        lt.style = 'Table Grid'
        head = lt.rows[0].cells
        for i, t in enumerate(['时间', '节点', '动作', '操作人', '意见']):
            head[i].text = t
        for log in payload.logs:
            cells = lt.add_row().cells
            cells[0].text = str(log.get('create_time') or log.get('createTime') or '-')
            cells[1].text = str(log.get('to_node_code') or log.get('toNodeCode') or '-')
            cells[2].text = str(log.get('action_name') or log.get('actionName') or '-')
            cells[3].text = str(log.get('operator_name') or log.get('operatorName') or '-')
            cells[4].text = str(log.get('opinion') or '')

    doc.add_heading('二、逐条意见现场核查记录', level=1)
    opinions = block.get('opinions') or []
    if not opinions:
        doc.add_paragraph('该项目尚未生成现场核查任务/照片，本单仅含审核意见。')
    for idx, op in enumerate(opinions, 1):
        doc.add_heading(f'意见{idx}：{op["title"]}', level=2)
        meta = f'风险等级：{op["risk_level"] or "-"}   阶段：{op["stage_name"] or "-"}'
        doc.add_paragraph(meta)
        p = doc.add_paragraph()
        p.add_run('意见内容：').bold = True
        p.add_run(op['content'])
        state = _OP_STATUS_LABEL.get(op['status'], op['status'])
        p = doc.add_paragraph()
        p.add_run('核查状态：').bold = True
        p.add_run(f'{state}    {op["check_time"] or ""}'.strip())
        if op.get('return_reason'):
            p = doc.add_paragraph()
            p.add_run('退回原因：').bold = True
            p.add_run(op['return_reason'])
        p = doc.add_paragraph()
        p.add_run('核查意见：').bold = True
        p.add_run(op.get('check_opinion') or '-')
        photos = op.get('photos') or []
        if photos:
            p = doc.add_paragraph()
            p.add_run(f'现场照片（{len(photos)}张）：').bold = True
            for path in photos:
                try:
                    doc.add_picture(str(path), width=Inches(3.4))
                except Exception:
                    doc.add_paragraph(f'[照片缺失] {path.name}')
        else:
            doc.add_paragraph('现场照片：暂无')

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@router.post("/export/workflow-word")
def export_workflow_word(payload: WorkflowWordExportPayload = Body(...)) -> Response:
    block = _patrol_check_block(payload.project_id)
    try:
        data = _build_workflow_word(payload, block)
    except Exception as exc:  # pragma: no cover - 生成失败兜底
        LOGGER.exception("Word 导出失败")
        raise HTTPException(status_code=500, detail=f"Word 生成失败：{exc}") from exc
    name = f"{payload.project_name or '项目'}-{payload.stage_name or '流程'}-终审意见与现场核查记录"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}.docx"},
    )
